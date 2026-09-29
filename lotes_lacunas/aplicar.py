"""Aplica os resultados dos 10 lotes de lacunas (vida util / IC / volume) a dados_fck.

2026-09-29. Regras:
  - so preenche celulas vazias (nunca apaga valor ja gravado), com as excepcoes
    em SOBREPOR (valores gravados que a leitura integral mostrou errados ou vagos);
  - tipo_corpo_prova so e substituido quando o valor antigo nao tem dimensao
    (sem digitos) e o novo tem;
  - o grupo IC (valor/tipo/n) so e escrito se o valor IC da linha estiver vazio;
  - FIGURAS: leituras de barras de erro feitas pessoalmente nos recortes
    renderizados (shrink-drive-pdf.yml), que os subagentes nao conseguiram ver;
  - EXTRA: desvios padrao impressos em tabelas que os subagentes registaram nas
    notas ("nao pedido") e que ainda nao estavam gravados.
Correr da raiz do repo; depois correr recalcular_ci.py.
"""
import glob
import json
import re

import openpyxl

BASE = 'meta_analysis_base.xlsx'
C = {'idade_dias': 10, 'tipo_corpo_prova': 11, 'vida_util_anos': 19, 'vida_util_citacao': 20,
     'vol_corpo_prova_m3': 23, 'vol_corpo_prova_fonte': 24, 'vol_corpo_prova_citacao': 25,
     'ic_fck_valor': 26, 'ic_fck_tipo': 27, 'ic_fck_n': 28, 'ic_fck_citacao': 29,
     'ic_co2_valor': 30, 'ic_co2_tipo': 31, 'ic_co2_n': 32, 'ic_co2_citacao': 33}

# (sheet_row, campo) -> valor que substitui o gravado
SOBREPOR = {
    (1691, 'idade_dias'): 28,   # 4696: fck gravado e o dos 28 d (Tabela), nao 1 d
    (935, 'idade_dias'): 28,    # 3225: 'nao informado' -> 28 d declarados no texto
}
for r in range(1958, 1986):     # 3203 VACLAVIK: fck gravados coincidem com os 28 d
    SOBREPOR[(r, 'idade_dias')] = 28
NAO_ESCREVER = {('4698', 'tipo_corpo_prova')}  # sem suporte no artigo

FIG = 'leitura pessoal das barras de erro na figura renderizada (2026-09-29)'
FIGURAS = {
    # HAIGH 4849, Fig.5 p.10, barras = DP (texto p.9)
    1821: ('fck', 1.0, 'desvio_padrao', None, 'Fig.5 p.10; ' + FIG),
    1822: ('fck', 1.0, 'desvio_padrao', None, 'Fig.5 p.10; ' + FIG),
    1823: ('fck', 1.5, 'desvio_padrao', None, 'Fig.5 p.10; barra assimetrica (+1/-2), media ' + FIG),
    # NASSAR 2755, Fig.5, 28 d, 3 provetes, DP
    187: ('fck', 5.0, 'desvio_padrao', 3, 'Fig.5 (28 d, Ref); ' + FIG + ', precisao ~1 MPa'),
    188: ('fck', 6.1, 'desvio_padrao', 3, 'Fig.5 (28 d, M1-URCF-5); ' + FIG + ', precisao ~1 MPa'),
    189: ('fck', 4.3, 'desvio_padrao', 3, 'Fig.5 (28 d, M5-URCF-25); ' + FIG + ', precisao ~1 MPa'),
    190: ('fck', 4.9, 'desvio_padrao', 3, 'Fig.5 (28 d, M6-CRCF-5); ' + FIG + ', precisao ~1 MPa'),
    191: ('fck', 4.0, 'desvio_padrao', 3, 'Fig.5 (28 d, M9-CRCF-20); ' + FIG + ', precisao ~1 MPa'),
    # SIRICO 2895, Fig.10 p.12 'Average compressive strength (and corresponding standard deviation)'
    786: ('fck', 2.6, 'desvio_padrao', 3, 'Fig.10 p.12 (28 d, V0); ' + FIG),
    787: ('fck', 1.7, 'desvio_padrao', 3, 'Fig.10 p.12 (28 d, V10); ' + FIG),
    788: ('fck', 0.7, 'desvio_padrao', 3, 'Fig.10 p.12 (28 d, V15); ' + FIG),
    789: ('fck', 0.8, 'desvio_padrao', 3, 'Fig.10 p.12 (28 d, V20); ' + FIG),
}

EXTRA = {
    # 3176, Table 4 (DP por mistura, n=3)
    1937: ('fck', 0.56, 'desvio_padrao', 3, 'Table 4: DP 5A = 0,56 MPa'),
    1938: ('fck', 0.39, 'desvio_padrao', 3, 'Table 4: DP 5R = 0,39 MPa'),
    1939: ('fck', 0.91, 'desvio_padrao', 3, 'Table 4: DP 10A = 0,91 MPa'),
    1940: ('fck', 1.16, 'desvio_padrao', 3, 'Table 4: DP 10R = 1,16 MPa'),
    1941: ('fck', 1.19, 'desvio_padrao', 3, 'Table 4: DP 15A = 1,19 MPa'),
    1942: ('fck', 0.96, 'desvio_padrao', 3, 'Table 4: DP 15R = 0,96 MPa'),
    1943: ('fck', 1.57, 'desvio_padrao', 3, 'Table 4: DP 20A = 1,57 MPa'),
    1944: ('fck', 0.91, 'desvio_padrao', 3, 'Table 4: DP 20R = 0,91 MPa'),
    1945: ('fck', 3.95, 'desvio_padrao', 3, 'Table 4: DP 30A = 3,95 MPa'),
    1946: ('fck', 0.60, 'desvio_padrao', 3, 'Table 4: DP 30R = 0,60 MPa'),
    1947: ('fck', 1.03, 'desvio_padrao', 3, 'Table 4: DP 40A = 1,03 MPa'),
    1948: ('fck', 0.66, 'desvio_padrao', 3, 'Table 4: DP 40R = 0,66 MPa'),
    # 3786, Tabela 3, Condicao 2 (n=3)
    2013: ('fck', 3.46, 'desvio_padrao', 3, 'Tabela 3: 26,81 (3,46), Condicao 2'),
    2014: ('fck', 3.24, 'desvio_padrao', 3, 'Tabela 3: 26,41 (3,24), Condicao 2'),
    2015: ('fck', 0.98, 'desvio_padrao', 3, 'Tabela 3: 21,54 (0,98), Condicao 2'),
    # 4505, Tabela 8 (4 ensaios por tipo de fibra)
    2068: ('fck', 0.96, 'desvio_padrao', 4, 'Tabela 8: DP CFG_12 MIN = 0,96 MPa'),
    2069: ('fck', 1.79, 'desvio_padrao', 4, 'Tabela 8: DP CFG_36 MIN = 1,79 MPa'),
    2070: ('fck', 0.44, 'desvio_padrao', 4, 'Tabela 8: DP CFP MIN = 0,44 MPa'),
    # 3309 DADD, Tab.2: CoV do fc-28 (n=6) -> DP = CoV x fc
    346: ('fck', 1.43, 'desvio_padrao', 6, 'Tab.2: CoV fc-28 VAC 0,030 x 47,61 MPa (DP derivado)'),
    1993: ('fck', 2.94, 'desvio_padrao', 6, 'Tab.2: CoV fc-28 RAI 0,053 x 55,5 MPa (DP derivado)'),
    1994: ('fck', 1.05, 'desvio_padrao', 6, 'Tab.2: CoV fc-28 RAII 0,022 x 47,91 MPa (DP derivado)'),
    1995: ('fck', 3.76, 'desvio_padrao', 6, 'Tab.2: CoV fc-28 RAIII 0,078 x 48,26 MPa (DP derivado)'),
    # 3281 NATH, Tabela 5: DP Monte Carlo do CO2
    332: ('co2', 13.1, 'desvio_padrao', None, 'Tabela 5: DP Monte Carlo FA0 13,1 kg (CV 3,8%)'),
    333: ('co2', 9.4, 'desvio_padrao', None, 'Tabela 5: DP Monte Carlo FA30 9,4 kg (CV 3,5%)'),
    334: ('co2', 9.2, 'desvio_padrao', None, 'Tabela 5: DP Monte Carlo FA40 9,2 kg (CV 3,42%)'),
}


def vazio(v):
    return v is None or (isinstance(v, str) and v.strip() == '')


def clean(s):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(s)) if isinstance(s, str) else s


def main():
    wb = openpyxl.load_workbook(BASE)
    ws = wb['dados_fck']
    cont = {}

    def put(r, campo, v):
        ws.cell(r, C[campo]).value = clean(v)
        cont[campo] = cont.get(campo, 0) + 1

    for f in sorted(glob.glob('lotes_lacunas/resultado_lote*.json')):
        for art in json.load(open(f)):
            for p in art['pares']:
                r = p['sheet_row']
                assert ws.cell(r, 1).value == art['row_id'], (f, r)
                ic_vazio = {g: vazio(ws.cell(r, C[f'ic_{g}_valor']).value) for g in ('fck', 'co2')}
                # 3894: intervalo min-max dado como lista -> semi-amplitude
                for g in ('fck', 'co2'):
                    v = p.get(f'ic_{g}_valor')
                    if isinstance(v, list):
                        p[f'ic_{g}_min'], p[f'ic_{g}_max'] = v
                        p[f'ic_{g}_valor'] = round((v[1] - v[0]) / 2, 2)
                    if p.get(f'ic_{g}_min') is not None and p.get(f'ic_{g}_max') is not None:
                        cit = p.get(f'ic_{g}_citacao') or ''
                        extra = f"min {p[f'ic_{g}_min']}; max {p[f'ic_{g}_max']}; valor = semi-amplitude"
                        if 'semi-amplitude' not in cit:
                            p[f'ic_{g}_citacao'] = (cit + ' | ' if cit else '') + extra
                for campo in C:
                    v = p.get(campo)
                    if (r, campo) in SOBREPOR:
                        put(r, campo, SOBREPOR[(r, campo)])
                        continue
                    if vazio(v) or (str(art['row_id']), campo) in NAO_ESCREVER:
                        continue
                    cur = ws.cell(r, C[campo]).value
                    if campo == 'tipo_corpo_prova':
                        if vazio(cur) or (not re.search(r'\d', str(cur)) and re.search(r'\d', str(v))):
                            put(r, campo, v)
                    elif campo == 'vol_corpo_prova_fonte':
                        if vazio(cur) or (cur == 'indisponivel' and not vazio(p.get('vol_corpo_prova_m3'))):
                            put(r, campo, v)
                    elif campo in ('ic_fck_tipo', 'ic_fck_n', 'ic_co2_tipo', 'ic_co2_n'):
                        g = campo.split('_')[1]
                        if ic_vazio[g] and (
                                vazio(cur) or not vazio(p.get(f'ic_{g}_valor'))):
                            put(r, campo, v)
                    elif vazio(cur):
                        put(r, campo, v)

    for dic in (FIGURAS, EXTRA):
        for r, (g, val, tipo, n, cit) in dic.items():
            if not vazio(ws.cell(r, C[f'ic_{g}_valor']).value):
                continue
            put(r, f'ic_{g}_valor', val)
            put(r, f'ic_{g}_tipo', tipo)
            if n is not None:
                put(r, f'ic_{g}_n', n)
            put(r, f'ic_{g}_citacao', cit)

    wb.save(BASE)
    for k in C:
        print(k, cont.get(k, 0))


if __name__ == '__main__':
    main()

"""F27 — TOSIC 4625: GWP da Tabela 4 corrigido de «g» para «kg» (decisao da Kie, 2026-10-16: «Pode corrigir g para kg (b)»).

Pergunta aberta C4 de metanalise/DECISOES_METANALISE.md; substitui a decisao de 2026-09-12 (manter o rotulo «g» tal
como impresso). Releitura pessoal do PDF a 2026-10-16 (texto completo + imagem da Tab.4, p.15). Prova, toda com dados
impressos no proprio artigo:
- UF = 1 m2 de laje com h = 0,2 m (0,2 m3 de betao, p.12); Tab.1 (p.13) da o cimento medio 340-381 kg/m3, ou seja
  68-76 kg de cimento por UF;
- o GWP da Tab.4 (68,5-74,7 por UF) acompanha esse teor quase 1:1 (NAC C25/30 68 kg -> 68,5; NAC C30/37 74,8 kg ->
  74,7): ~1 kg CO2 por kg de cimento, o esperado com as EPD CEMBUREAU que o artigo diz usar (Tab.3);
- em gramas daria 1 g CO2 por kg de cimento, fisicamente impossivel (so a descarbonatacao liberta ~500 g/kg de
  clinquer), e o GWP ficaria abaixo da acidificacao da mesma tabela (AP 155,8 g SO2/UF).
Precedentes de gralha corrigida com dados do proprio artigo: 2720 LI («1 L» -> 1 m3) e KIM (F17, /10).
"""
import re

import openpyxl

META = 'meta_analysis_base.xlsx'
HOJE = '2026-10-16'
# linha de dados_fck -> GWP impresso por UF (Tab.4 p.15)
GWP = {1659: 68.5, 1660: 69.6, 1661: 69.2, 1662: 74.7, 1663: 74.6, 1664: 72.9}
UF_M3 = 0.2
ALERTA = ('[ALERTA ETAPA4: rotulo «g CO2-eq» da Tab.4 corrigido para kg (F27, Kie 2026-10-16) - prova no proprio '
          'artigo: cimento da Tab.1 (340-381 kg/m3) x UF 0,2 m3 = 68-76 kg, GWP impresso 68,5-74,7 (~1 kg CO2/kg '
          'cimento, EPD CEMBUREAU); em g seria fisicamente impossivel; confianca baixa]')


def clean(s):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(s))


def main():
    wb = openpyxl.load_workbook(META)
    ws = wb['dados_fck']
    for r, g in GWP.items():
        assert ws.cell(r, 1).value == 4625, r
        antes = ws.cell(r, 37).value
        novo = round(g / UF_M3, 2)
        ws.cell(r, 37).value = novo
        ws.cell(r, 14).value = novo
        ws.cell(r, 15).value = clean(f'kg CO2-eq/m3 [CORRIGIDO {HOJE}, F27: Tab.4 impressa em «g CO2-eq» por UF, lida '
                                     f'como kg; {g} kg / 0,2 m3] {ALERTA}')
        ws.cell(r, 36).value = clean(f'correcao {HOJE} (F27): {g} kg (rotulo «g» da Tab.4 lido como kg) / 0,2 m3 (UF) '
                                     f'= {novo} kg CO2-eq/m3. {ALERTA}')
        ws.cell(r, 13).value = clean(f'{ws.cell(r, 13).value or ""} [CO2 {antes} -> {novo} {HOJE} (F27, Kie: «Pode '
                                     f'corrigir g para kg (b)»; substitui a decisao de 2026-09-12) {ALERTA}]')
    wb.save(META)
    print('corrigidas', len(GWP))


if __name__ == '__main__':
    main()

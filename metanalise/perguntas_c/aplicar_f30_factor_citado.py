"""F30 — factor generico de EPD/base de dados citado = Aceito (decisao da Kie, 2026-10-16: «Fator generico de EPD ou
base de dados sem ACV entram sim, desde que descritos e citados»).

Supera o fundamento «factor generico sem ACV propria» da F17 (JATIVA), F20 (CHOI) e F29 (OH). Continuam rejeitados os
casos em que o factor nao tem fonte (3917 UNGUREANU, que alem disso exclui o betao do calculo) ou em que o CO2 nao se
leva ao betao (3691 MITOULIS, F18 A1).

- 4930 OH: factor da base de EPD coreana (ref.[10]). Na Tab.2 (p.3) as linhas Concrete 3.532 / Rebar 259 kg/m3 estao
  trocadas: a Fig.9(a) (p.5, comparacao por volume) da RC = 2,3% = 259 / 11.241 do aco e madeira 0,78% = 87,64 / 11.241.
  Par gravado com 259 kg CO2-eq/m3 (C30, fck 30), com alerta.
- 4620 CHOI: factor ICE v2.0 0,141 kg/kg (p.4), citado; sem massa volumica impressa -> Aceito sem converter (F18).
- 3069 JATIVA: factor de literatura 0,86 t CO2/t cimento (ref.[3]), citado -> Aceito; continua sem par convertivel.
"""
import re

import openpyxl

ART = 'articles_final.xlsx'
META = 'meta_analysis_base.xlsx'
HOJE = '2026-10-16'
LIN = {
    4930: 'factor da base de EPD coreana citado (ref.[10]); par com 259 kg/m3 (linhas Concrete/Rebar da Tab.2 trocadas, '
          'prova na Fig.9(a): RC = 2,3% = 259/11.241)',
    4620: 'factor ICE v2.0 citado (p.4); sem massa volumica impressa -> Aceito sem converter (F18)',
    3069: 'factor de literatura citado (ref.[3], 0,86 t CO2/t cimento) -> Aceito; sem par convertivel',
}
OH_DADOS = 1835
ALERTA = ('[ALERTA ETAPA4: factor generico da base de EPD coreana (ref.[10]), nao ACV propria; na Tab.2 as linhas Concrete '
          '3.532 / Rebar 259 estao trocadas - a Fig.9(a) da RC = 2,3% = 259/11.241 do aco (F30, Kie 2026-10-16)]')


def clean(s):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(s))


def main():
    wa = openpyxl.load_workbook(ART)
    ws = wa['Articles']
    for r, m in LIN.items():
        assert ws.cell(r, 20).value == 'Rejeitado' and ws.cell(r, 40).value == 'Rejeitado', r
        assert ws.cell(r, 39).value == 'Aceito', r
        ws.cell(r, 20).value = 'Elegível'
        ws.cell(r, 40).value = 'Aceito'
        ws.cell(r, 43).value = None
        ws.cell(r, 44).value = None
        ws.cell(r, 32).value = clean(f'{ws.cell(r, 32).value or ""} [F30 {HOJE} (Kie: «Fator generico de EPD ou base de '
                                     f'dados sem ACV entram sim, desde que descritos e citados»): E2_co2 Aceito - {m}]')
    wa.save(ART)

    wb = openpyxl.load_workbook(META)
    wd = wb['dados_fck']
    assert wd.cell(OH_DADOS, 1).value == 4930
    wd.cell(OH_DADOS, 37).value = 259.0
    wd.cell(OH_DADOS, 14).value = 259.0
    wd.cell(OH_DADOS, 15).value = clean(f'kg CO2-eq/m3 concreto {ALERTA}')
    wd.cell(OH_DADOS, 36).value = clean(f'correcao {HOJE} (F30): 259 kg CO2-eq/m3 = linha «Rebar» da Tab.2, que e a do '
                                        f'betao (troca provada pela Fig.9(a)); nativo por m3. {ALERTA}')
    wd.cell(OH_DADOS, 13).value = clean(f'{wd.cell(OH_DADOS, 13).value or ""} [CO2 NAO CONVERTIDO -> 259 {HOJE} (F30, '
                                        f'verificacao pessoal) {ALERTA}]')
    wb.save(META)
    print('repostos', len(LIN), '| OH co2 = 259')


if __name__ == '__main__':
    main()

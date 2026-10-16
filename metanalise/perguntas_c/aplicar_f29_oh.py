"""F29 — OH 4930 rejeitado em E2_co2 (decisao da Kie, 2026-10-16: «Pode rejeitar em E2_co2 (b)»).

Pergunta aberta C5 de metanalise/DECISOES_METANALISE.md. Releitura pessoal do PDF a 2026-10-16 (6 paginas, texto + imagem
da Fig.9, p.5):
- a Tab.2 (p.3) da os factores de emissao por m3 tirados da base de EPD do Ministerio do Ambiente coreano (ref.[10]):
  GLT 87,64; aco 11.241; «Concrete 3.532»; «Rebar 259»; as linhas do betao e da armadura estao trocadas - a Fig.9(a)
  (comparacao por volume, normalizada ao aco) da RC = 2,3% = 259 / 11.241 e madeira 0,78% = 87,64 / 11.241;
- mesmo corrigida a troca, o CO2 do betao e um factor generico de base de EPD (a classe e a EPD nao sao identificadas),
  multiplicado pelas quantidades de uma viga; o resultado proprio do artigo e por metro de viga, betao + armadura, sem
  volume impresso. Precedente F20 (CHOI, factor ICE) e JATIVA/UNGUREANU: factor generico sem ACV propria do betao.
"""
import re

import openpyxl

ART = 'articles_final.xlsx'
META = 'meta_analysis_base.xlsx'
HOJE = '2026-10-16'
ROW = 4930
MOTIVO = ('precedente F20: o unico CO2 do betao e um factor generico da base de EPD coreana (ref.[10], classe e EPD nao '
          'identificadas), sem ACV propria do betao; resultado proprio so por metro de viga (betao + armadura), sem '
          'volume impresso. Nota: na Tab.2 (p.3) as linhas Concrete 3.532 / Rebar 259 kg/m3 estao trocadas - a Fig.9(a) '
          'da RC = 2,3% = 259/11.241 do aco')


def clean(s):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(s))


def main():
    wa = openpyxl.load_workbook(ART)
    ws = wa['Articles']
    assert ws.cell(ROW, 20).value == 'Elegível' and ws.cell(ROW, 40).value == 'Aceito'
    ws.cell(ROW, 20).value = 'Rejeitado'
    ws.cell(ROW, 40).value = 'Rejeitado'
    ws.cell(ROW, 43).value = 2
    ws.cell(ROW, 44).value = 'E2_co2'
    ws.cell(ROW, 32).value = clean(f'{ws.cell(ROW, 32).value or ""} [F29 {HOJE} (Kie: «Pode rejeitar em E2_co2 (b)»): '
                                   f'E2_co2 Rejeitado - {MOTIVO}]'.strip())
    wa.save(ART)

    wb = openpyxl.load_workbook(META)
    wd = wb['dados_fck']
    n = 0
    for r in range(2, wd.max_row + 1):
        if wd.cell(r, 1).value != ROW or wd.cell(r, 37).value is None:
            continue
        antes = wd.cell(r, 37).value
        wd.cell(r, 37).value = None
        wd.cell(r, 44).value = None
        wd.cell(r, 36).value = clean(f'NAO CONVERTIDO - F29 {HOJE}: {MOTIVO}')[:1500]
        wd.cell(r, 13).value = clean(f'{wd.cell(r, 13).value or ""} [CO2 {antes} -> NAO CONVERTIDO {HOJE} (F29, '
                                     f'verificacao pessoal): {MOTIVO}]')
        n += 1
    wb.save(META)
    print('rejeitado', ROW, '| pares retirados', n)


if __name__ == '__main__':
    main()

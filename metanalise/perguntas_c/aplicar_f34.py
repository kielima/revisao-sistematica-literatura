"""F34 - Pergunta C4: mesmos numeros em artigos diferentes (Kie, 2026-10-17: «Pode seguir a recomendacao nos dois»).

- 2854 FILHO: Duplicado da 2744 (feito antes, ver DECISOES F34).
- 4027 TEH & WIEDMANN 2018 (J. Econ. Struct. 7:15): decomposicao do mesmo inventario hibrido da 4026 (mesmos dados de
  Teh et al. 2017); os totais (517 / 363 kg CO2/m3 de RCA OPC/GPC) coincidem por construcao -> Duplicado da 4026.
- 2695 JANG 2022 (Buildings 12:51), linha 655 de dados_fck (24 MPa, 414 kg/m3): o mesmo factor KEITI/LCI DB coreano que
  a 2694 usa; o par e contado uma so vez, na 2694 -> o par sai da 2695 (o artigo continua Elegivel: outros dados/contexto).
"""
import re

import openpyxl

HOJE = '2026-10-17'


def clean(s):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(s))


wa = openpyxl.load_workbook('articles_final.xlsx')
ws = wa['Articles']
assert ws.cell(4027, 20).value == 'Elegível'
ws.cell(4027, 20).value = 'Duplicado'
ws.cell(4027, 32).value = clean(f'{ws.cell(4027, 32).value or ""} [DUPLICADO {HOJE}, decisao da Kie (pergunta C4): '
    'decomposicao do mesmo inventario hibrido da linha 4026 (mesmos dados de Teh et al. 2017; os totais 517/363 '
    'kg CO2/m3 da Fig.8 coincidem por construcao); a 4026 fica como canonica. Colunas E1-E4 mantidas como historico.]')
wa.save('articles_final.xlsx')

wb = openpyxl.load_workbook('meta_analysis_base.xlsx')
wd = wb['dados_fck']
assert wd.cell(655, 1).value == 2695, wd.cell(655, 1).value
antes = wd.cell(655, 37).value
motivo = ('mesmo factor KEITI/LCI DB coreano (24 MPa, 414 kg CO2/m3) que a 2694 ja usa: o par conta-se uma so vez '
          '(F34, Kie 2026-10-17)')
wd.cell(655, 37).value = None
wd.cell(655, 44).value = None
wd.cell(655, 36).value = clean(f'NAO CONVERTIDO - F34 {HOJE}: {motivo}')[:1500]
wd.cell(655, 13).value = clean(f'{wd.cell(655, 13).value or ""} [CO2 {antes} -> retirado {HOJE}: {motivo}]')
wb.save('meta_analysis_base.xlsx')
print('ok', antes)

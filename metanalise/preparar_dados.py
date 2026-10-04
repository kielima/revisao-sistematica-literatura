"""Prepara o ficheiro de entrada da metanalise estatistica (rma.mv REML).

Selecciona exactamente os mesmos pares do painel kielima.com/metanalise (usa as
funcoes e filtros do metanalise-export.py) e acrescenta, por par, o que o modelo
precisa: estudo, observacao, ci, ci_mpa, estrato E4, familia, base Portland,
idade, volume, vida util e a dispersao publicada (IC) do fck e do CO2.

Saida: metanalise/dados_metanalise.csv (UTF-8). Correr da raiz do repo:
    python3 metanalise/preparar_dados.py
Decisoes metodologicas: metanalise/DECISOES_METANALISE.md.
"""
import csv
import importlib.util
import os

import openpyxl

spec = importlib.util.spec_from_file_location('mexp', 'metanalise-export.py')
mexp = importlib.util.module_from_spec(spec)
spec.loader.exec_module(mexp)

SAIDA = 'metanalise/dados_metanalise.csv'
COLS = ['obs_id', 'estudo', 'rotulo', 'material', 'fck', 'co2_kg_m3', 'ci', 'estrato', 'familia', 'base_portland',
        'idade_dias', 'vol_m3', 'ci_mpa', 'vida_util', 'ic_fck_valor', 'ic_fck_tipo', 'ic_fck_n',
        'ic_co2_valor', 'ic_co2_tipo', 'ic_co2_n']


def main():
    status = {}
    af = openpyxl.load_workbook('articles_final.xlsx', read_only=True)['Articles']
    for i, row in enumerate(af.iter_rows(min_row=2, max_col=20, values_only=True), start=2):
        status[i] = row[19]
    ws = openpyxl.load_workbook(mexp.BASE, read_only=True)['dados_fck']
    linhas = []
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        g = lambda c: row[c - 1] if c - 1 < len(row) else None
        row_id, fck, co2 = g(1), mexp.tofloat(g(8)), mexp.tofloat(g(14))
        co2m3, estrato = mexp.tofloat(g(37)), g(22)
        if fck in (None, 0) or co2 is None or status.get(row_id) != 'Elegível':
            continue
        if estrato not in mexp.TIERCODE or co2m3 is None:
            continue
        if not (mexp.FCK_MIN <= fck <= mexp.FCK_MAX and mexp.CO2_MIN <= co2m3 <= mexp.CO2_MAX):
            continue
        vol = mexp.tofloat(g(23))
        vida = mexp.tofloat(g(19))
        ci_mpa = co2m3 * vol / fck if vol else None
        linhas.append({
            'obs_id': r, 'estudo': row_id, 'rotulo': mexp.short_label(g(4), g(5)),
            'material': str(g(7) or '')[:120], 'fck': fck, 'co2_kg_m3': co2m3, 'ci': co2m3 / fck,
            'estrato': estrato, 'familia': mexp.classify(row_id, g(7), g(3)),
            'base_portland': g(40) or 'misto', 'idade_dias': mexp.tofloat(g(10)), 'vol_m3': vol,
            'ci_mpa': ci_mpa, 'vida_util': vida,
            'ic_fck_valor': mexp.tofloat(g(26)), 'ic_fck_tipo': g(27), 'ic_fck_n': mexp.tofloat(g(28)),
            'ic_co2_valor': mexp.tofloat(g(30)), 'ic_co2_tipo': g(31), 'ic_co2_n': mexp.tofloat(g(32)),
        })
    os.makedirs('metanalise', exist_ok=True)
    with open(SAIDA, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=COLS)
        w.writeheader()
        w.writerows(linhas)
    print(f'{len(linhas)} pares de {len({l["estudo"] for l in linhas})} estudos -> {SAIDA}')


if __name__ == '__main__':
    main()

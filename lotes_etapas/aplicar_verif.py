"""F39: aplica os vereditos da verificacao completa (verif_resultado*.json) a dados_fck.
Politica (2026-11-01): retirar = unidade nao e m3 de betao (argamassa, pasta, madeira, m2...) ou leitura sem rotulos;
corrigir = so colunas de etapa; ignora co2_kg_m3 (valor do banco -> achado para a E2), ignora A4=0 por fronteira A1-A3 (so
com declaracao textual: PERMITE_ZERO); KINA a2=0 fica (artigo declara transporte nao considerado). Idempotente.
Uso: python3 lotes_etapas/aplicar_verif.py [--gravar]"""
import glob, json, sys, collections
import openpyxl
N = ['co2_a1','co2_a2','co2_a3','co2_a4','co2_a1a3_agregado','etapas_tipo','soma_confere','etapas_local','etapas_citacao','etapas_nota']
STAGE = {'a1':'co2_a1','a2':'co2_a2','a3':'co2_a3','a4':'co2_a4','a1a3':'co2_a1a3_agregado'}
PERMITE_ZERO = {4418, 4993, 3321}          # A4/A2 = 0 com declaracao textual no artigo
SALTAR_ROWS = {(3071,300),(3071,303)}       # Tab.4 incoerente com Tab.A5/A6
RETIRAR_ESTUDOS = {4476: 'unidade m3 de argamassa alcali-activada (FARIDMEHR)', 4820: 'unidade m3 de madeira lamelada colada (ASCIONE)'}
LIMPAR_CAMPOS = {(4675,'co2_a1a3_agregado'): 'fronteira inclui o transporte ate ao local final (A1-A4), nao A1-A3 (IMTIAZ)'}
gravar = '--gravar' in sys.argv
wb = openpyxl.load_workbook('meta_analysis_base.xlsx'); ws = wb['dados_fck']
h = [c.value for c in ws[1]]; c0 = h.index('co2_a1') + 1
col = lambda n: c0 + N.index(n)
def nota(r, t):
    c = ws.cell(r, col('etapas_nota')); 
    if t not in (c.value or ''): c.value = ((c.value or '') + ' ' + t).strip()
def limpar(r, motivo):
    for n in ('co2_a1','co2_a2','co2_a3','co2_a4','co2_a1a3_agregado'): ws.cell(r, col(n)).value = None
    ws.cell(r, col('etapas_tipo')).value = 'retirado'; nota(r, f'[retirado 2026-11-01: {motivo}]')
cont = collections.Counter(); ignorados = []
ver = {}
for f in sorted(glob.glob('lotes_etapas/verif_resultado*.json')):
    for e in json.load(open(f)):
        for p in e['pares']: ver[(e['row_id'], p['sheet_row'])] = p
for (rid, r), p in ver.items():
    assert ws.cell(r, 1).value == rid, (rid, r)
    v = p['veredito']
    if v == 'retirar' or rid in RETIRAR_ESTUDOS:
        limpar(r, (RETIRAR_ESTUDOS.get(rid) or (p.get('motivo') or '')[:160])); cont['retirar'] += 1; continue
    if v == 'corrigir' and (rid, r) not in SALTAR_ROWS:
        vc = p.get('valor_correto')
        if not isinstance(vc, dict): ignorados.append((rid, r, 'valor_correto nao e dict')); continue
        for k, val in vc.items():
            if k not in STAGE: ignorados.append((rid, r, f'{k} (nao e coluna de etapa)')); continue
            if val == 0 and rid not in PERMITE_ZERO: ignorados.append((rid, r, f'{k}=0 sem declaracao textual')); continue
            ws.cell(r, col({'a1':'co2_a1','a2':'co2_a2','a3':'co2_a3','a4':'co2_a4','a1a3':'co2_a1a3_agregado'}[k])).value = val
            cont['corrigir'] += 1
        nota(r, '[verificado 2026-11-01: etapa corrigida]')
# estudos com unidade nao-betao sem veredito de retirar explicito
for r in range(2, ws.max_row + 1):
    rid = ws.cell(r, 1).value
    if rid in RETIRAR_ESTUDOS and ws.cell(r, col('etapas_tipo')).value not in (None, '', 'retirado'):
        limpar(r, RETIRAR_ESTUDOS[rid]); cont['retirar_estudo'] += 1
    for (i, k), mot in LIMPAR_CAMPOS.items():
        if rid == i and ws.cell(r, col(k)).value not in (None, ''):
            ws.cell(r, col(k)).value = None; nota(r, f'[retirado {k}: {mot}]'); cont['limpar_campo'] += 1
    # restaurar A4 do PUSHKAR 4297 (factor e distancia impressos no artigo, verificado)
    if rid == 4297 and r == 1264: ws.cell(r, col('co2_a4')).value = 47.0
print(dict(cont)); print('ignorados', len(ignorados)); [print(' ', x) for x in ignorados[:20]]
if gravar: wb.save('meta_analysis_base.xlsx'); print('gravado')

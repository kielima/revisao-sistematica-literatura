"""Cria a aba `estudos_secundarios` em articles_final.xlsx (2026-10-18, Kie: «pode avancar com a tabela de estudos
secundarios»). Uma linha por estudo citado por uma revisao (Revisao/Review sem ACV propria) cujos dados ainda
nao foram extraidos. Fontes: lotes_etapa2/{ONSONGO_2025,NASIR_2024}_ESTUDOS_SECUNDARIOS.csv e SCOPE_2021_ESTUDOS_LCA_TAB_A1.csv.

Colunas: id_sec | row_id_revisao | revisao | codigo_citacao | autor_ano | referencia_completa | doi | metodo |
observacao | row_id_banco (linha de `Articles` com o mesmo DOI, se existir) | status_banco | em_metanalise (1 = ja e
Elegivel/primario) | estado_extracao (pendente | ja_no_banco | sem_doi) | data.
Idempotente: recria a aba. Uso: python3 lotes_etapa2/criar_estudos_secundarios.py"""
import csv
import re

import openpyxl

HOJE = '2026-10-18'
REV = [(3001, 'ONSONGO et al., 2025', 'lotes_etapa2/ONSONGO_2025_ESTUDOS_SECUNDARIOS.csv'),
       (3163, 'NASIR et al., 2024', 'lotes_etapa2/NASIR_2024_ESTUDOS_SECUNDARIOS.csv'),
       (3671, 'SCOPE et al., 2021', 'lotes_etapa2/SCOPE_2021_ESTUDOS_LCA_TAB_A1.csv')]
COLS = ['id_sec', 'row_id_revisao', 'revisao', 'codigo_citacao', 'autor_ano', 'referencia_completa', 'doi', 'metodo',
        'observacao', 'row_id_banco', 'status_banco', 'em_metanalise', 'estado_extracao', 'data']


def nd(s):
    m = re.search(r'10\.\d{4,9}/\S+', str(s or '').lower())
    return m.group(0).rstrip('.,;') if m else None


def clean(s):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(s or ''))


wb = openpyxl.load_workbook('articles_final.xlsx')
ws = wb['Articles']
h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
por_doi = {}
for r in range(2, ws.max_row + 1):
    d = nd(ws.cell(r, h['doi']).value) or nd(ws.cell(r, h['url']).value)
    if d:
        por_doi.setdefault(d, []).append(r)
if 'estudos_secundarios' in wb.sheetnames:
    del wb['estudos_secundarios']
wo = wb.create_sheet('estudos_secundarios')
wo.append(COLS)
n = 0
for rid, nome, path in REV:
    for r in csv.DictReader(open(path, encoding='utf-8')):
        n += 1
        d = nd(r.get('doi'))
        # prefere a linha viva (Elegivel) quando o DOI esta em varias
        linhas = por_doi.get(d, []) if d else []
        linhas = [x for x in linhas if x != rid]
        viva = [x for x in linhas if ws.cell(x, h['status_triagem']).value == 'Elegível']
        rb = (viva or linhas or [None])[0]
        st = ws.cell(rb, h['status_triagem']).value if rb else None
        estado = 'sem_doi' if not d else ('ja_no_banco' if rb else 'pendente')
        wo.append([n, rid, nome, clean(r.get('codigo_citacao')), clean(r.get('autor_ano')), clean(r.get('referencia_completa')),
                   d, clean(r.get('metodo')), clean(r.get('observacao')), rb, st, 1 if st == 'Elegível' else 0, estado, HOJE])
wb.save('articles_final.xlsx')
import collections
c = collections.Counter((x[12]) for x in wo.iter_rows(min_row=2, values_only=True))
print(n, 'estudos secundarios;', dict(c), '; em metanalise:', sum(1 for x in wo.iter_rows(min_row=2, values_only=True) if x[11]))

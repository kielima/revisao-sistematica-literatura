"""F14 (Kie, 2026-10-11: «Pode corrigir o 3222 e o 4511»).

3222 ZHANG 2025: a linha legada (folha dados_fck, linha 933) juntava as 5 misturas
MT0-MT100 numa so celula de fck e de CO2, com co2_kg_m3 = 0 -- o exportador nao a
lia e o artigo ficava sem par no painel. Os valores ja tinham sido verificados no
PDF a 2026-09-26 (fck: Fig.10a por geometria vectorial, = 107,15 x (1+%) do texto
sec.5.3; CO2: Tabela 5 (traco) x Tabela 7 (GWP), conferido com a Fig.18 e com a
reducao de 1,11 kg/m3 do texto). Aqui so se desdobra a linha numa por mistura,
mesmo procedimento das linhas legadas multi-traco de 2026-09-28.
"""
import re
import openpyxl

BASE = 'meta_analysis_base.xlsx'
LINHA = 933
MIX = [('MT0', 'controlo, 0% de rejeito de molibdenio'), ('MT25', '25% de rejeito de molibdenio'),
       ('MT50', '50% de rejeito de molibdenio'), ('MT75', '75% de rejeito de molibdenio'),
       ('MT100', '100% de rejeito de molibdenio')]


def clean(s):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(s))


def main():
    wb = openpyxl.load_workbook(BASE)
    ws = wb['dados_fck']
    assert ws.cell(LINHA, 1).value == 3222
    fck_txt, co2_txt = str(ws.cell(LINHA, 8).value), str(ws.cell(LINHA, 14).value)
    fck = dict(re.findall(r'(MT\d+)(?:\(controlo\))?:([\d.]+)', fck_txt))
    co2 = dict(re.findall(r'(MT\d+):([\d.]+)', co2_txt))
    assert len(fck) == 5 and len(co2) == 5, (fck, co2)
    ncol = ws.max_column
    modelo = [ws.cell(LINHA, c).value for c in range(1, ncol + 1)]
    destino = [LINHA] + [ws.max_row + i for i in range(1, 5)]
    for (cod, desc), r in zip(MIX, destino):
        for c in range(1, ncol + 1):
            ws.cell(r, c).value = modelo[c - 1]
        ws.cell(r, 7).value = clean(f'UHPC com rejeito de mineracao de molibdenio (tailings), mistura {cod} ({desc})')
        ws.cell(r, 8).value = float(fck[cod])
        ws.cell(r, 14).value = float(co2[cod])
        ws.cell(r, 34).value = f'{co2[cod]} kg CO2/m3 ({cod})'
        ws.cell(r, 36).value = 'calculado pelo artigo e conferido: Tabela 5 (traco kg/m3) x Tabela 7 (GWP kg CO2/kg), p.5 e p.9; confere com a Fig.18'
        ws.cell(r, 37).value = float(co2[cod])
        obs = str(modelo[12] or '')
        ws.cell(r, 13).value = clean(f'{obs} [F14 2026-10-11: linha legada das 5 misturas desdobrada; esta linha = {cod}, '
                                     f'fck {fck[cod]} MPa, CO2 {co2[cod]} kg/m3 (antes co2_kg_m3 = 0)]')
    wb.save(BASE)
    print('linhas', destino)


if __name__ == '__main__':
    main()

"""Aplica a vida útil do ci_t a todos os pares do painel (F33, Kiê 2026-10-17).

Hierarquia aprovada pela Kiê («Aprovo essa ordem, média por família só como sensibilidade»):
  1. declarada pelo próprio estudo  -> dados_fck col 19 `vida_util_anos` (premissa própria; não muda, o QD depende dela)
  2. citada pelo estudo              -> resultado dos lotes (`citadas`, a preferida, só valor numérico único)
  3. norma conforme a aplicação      -> EN 1990 Tab. 2.1 / NBR 15575-1 Tab. C.5 (metanalise/NORMAS_VIDA_UTIL_REFERENCIAS.md)
  4. sem aplicação identificável     -> 50 anos (EN 1990 cat. 4)

Colunas novas em dados_fck (escritas só nos pares do painel, isto é, em dados_metanalise.csv):
  52 vida_util_ci_t     anos usados no ci_t principal (hierarquia)
  53 vida_util_fonte    declarada | citada | norma_aplicacao | norma_padrao
  54 vida_util_base     modelo | norma | literatura | premissa_sem_fonte (de onde vem o número)
  55 vida_util_citacao  local + frase do artigo, ou a regra da norma aplicada
  56 aplicacao          edificio | ponte | infraestrutura | pavimento | pre_fabricado_corrente | nao_definida
  57 vida_util_modelo   valor de base `modelo` (passos 1–2), para o cenário «família — modelos»
  58 vida_util_norma    valor de base `norma` (passos 1–2) ou, sem ele, o da norma pela aplicação (passos 3–4)
  59 vida_util_estudo   valor vindo do estudo (passos 1–2, qualquer base), para o cenário «média por família»

Uso: python3 lotes_vidautil/aplicar.py [--seco]
"""
import glob
import json
import re
import sys

import openpyxl

BASE = 'meta_analysis_base.xlsx'
NOVAS = {52: 'vida_util_ci_t', 53: 'vida_util_fonte', 54: 'vida_util_base', 55: 'vida_util_citacao',
         56: 'aplicacao', 57: 'vida_util_modelo', 58: 'vida_util_norma', 59: 'vida_util_estudo'}
NORMA_APLIC = {
    'edificio': (50, 'EN 1990 Tab. 2.1 cat. 4 (estruturas de edificios); NBR 15575-1 Tab. C.5 (estrutura, minimo 50)'),
    'pre_fabricado_corrente': (50, 'EN 1990 Tab. 2.1 cat. 4 (estruturas de edificios e correntes)'),
    'pavimento': (50, 'EN 1990 Tab. 2.1 cat. 4 (outras estruturas correntes) - por confirmar'),
    'ponte': (100, 'EN 1990 Tab. 2.1 cat. 5 (pontes e outras obras de engenharia civil)'),
    'infraestrutura': (100, 'EN 1990 Tab. 2.1 cat. 5 (pontes e outras obras de engenharia civil)'),
}
PADRAO = (50, 'EN 1990 Tab. 2.1 cat. 4 - aplicacao nao definida no estudo, vida util atribuida por nos')
BASES = {'modelo', 'norma', 'literatura', 'premissa_sem_fonte'}


def clean(s):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(s)) if s is not None else None


def num(v):
    try:
        f = float(v)
        return f if f > 0 else None
    except (TypeError, ValueError):
        return None


def carregar_resultados():
    res = {}
    for f in sorted(glob.glob('lotes_vidautil/resultado_lote*.json')):
        for r in json.load(open(f, encoding='utf-8')):
            res[int(r['row_id'])] = r
    return res


def citada_preferida(r):
    cits = [c for c in (r.get('citadas') or []) if num(c.get('valor'))]
    if not cits:
        return None
    pref = [c for c in cits if c.get('preferido')]
    return (pref or cits)[0]


def main():
    seco = '--seco' in sys.argv
    wl = json.load(open('lotes_vidautil/WORKLIST_2026-10-17.json', encoding='utf-8'))
    res = carregar_resultados()
    wb = openpyxl.load_workbook(BASE)
    ws = wb['dados_fck']
    for c, nome in NOVAS.items():
        ws.cell(1, c).value = nome
    conta = {'declarada': 0, 'citada': 0, 'norma_aplicacao': 0, 'norma_padrao': 0}
    alertas, faltam = [], []
    for est in wl:
        rid = est['row_id']
        r = res.get(rid)
        if r is None:
            faltam.append(rid)
            continue
        aplic = r.get('aplicacao') if r.get('aplicacao') in NORMA_APLIC or r.get('aplicacao') == 'nao_definida' else 'nao_definida'
        decl = r.get('declarada') or {}
        por_par = {int(p['sheet_row']): p for p in (r.get('declarada_por_par') or []) if p.get('sheet_row')}
        cit = citada_preferida(r)
        for p in est['pares']:
            row = p['sheet_row']
            v19 = num(ws.cell(row, 19).value)
            dp = por_par.get(row) or decl
            base_decl = dp.get('base') if dp.get('base') in BASES else None
            if v19:
                if not base_decl:
                    alertas.append(f'{rid} l.{row}: vida gravada {v19} mas o lote nao a encontrou como declarada')
                elif num(dp.get('valor')) and abs(num(dp.get('valor')) - v19) > 0.5:
                    alertas.append(f'{rid} l.{row}: vida gravada {v19} != lote {dp.get("valor")}')
                vida, fonte = v19, 'declarada'
                base = base_decl or 'premissa_sem_fonte'
                citacao = dp.get('citacao') or ws.cell(row, 20).value
            elif num(dp.get('valor')) and base_decl and dp.get('sheet_rows') != []:
                # o estudo declara a vida para este par mas a col 19 estava vazia: conta como declarada no ci_t
                # (a col 19 e o QD nao mudam; por confirmar, ver DECISOES F33)
                alertas.append(f'{rid} l.{row}: declarada {dp.get("valor")} ({base_decl}) usada no ci_t, col 19 vazia')
                vida, fonte, base = num(dp.get('valor')), 'declarada', base_decl
                citacao = dp.get('citacao')
            else:
                if cit:
                    vida, fonte = num(cit['valor']), 'citada'
                    base = cit.get('base') if cit.get('base') in BASES else 'literatura'
                    citacao = f"{cit.get('citacao') or ''} [fonte: {cit.get('fonte') or ''}]"
                elif aplic in NORMA_APLIC:
                    vida, citacao = NORMA_APLIC[aplic]
                    fonte, base = 'norma_aplicacao', 'norma'
                    citacao = f'{citacao}; aplicacao: {r.get("aplicacao_citacao") or ""}'
                else:
                    vida, citacao = PADRAO
                    fonte, base = 'norma_padrao', 'norma'
            conta[fonte] += 1
            do_estudo = vida if fonte in ('declarada', 'citada') else None
            modelo = vida if do_estudo and base == 'modelo' else None
            if do_estudo and base == 'norma':
                norma = vida
            else:
                norma = (NORMA_APLIC.get(aplic) or PADRAO)[0]
            valores = {52: vida, 53: fonte, 54: base, 55: clean(citacao)[:900] if citacao else None,
                       56: aplic, 57: modelo, 58: norma, 59: do_estudo}
            for c, v in valores.items():
                ws.cell(row, c).value = v
    print('pares por fonte:', conta)
    print('estudos sem resultado:', len(faltam), faltam[:20])
    print(f'{len(alertas)} alertas:')
    for a in alertas:
        print('  ', a)
    if not seco:
        wb.save(BASE)
        print('gravado', BASE)


if __name__ == '__main__':
    main()

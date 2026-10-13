"""Pais do autor de correspondencia dos artigos Elegiveis (decisao da Kie, 2026-10-13).

Kie: «So o pais do autor de correspondencia na coluna affiliation. BDTD Brasil, Crossref, os 70 que faltam
abrir o pdf e ver a correspondencia do primeiro autor.»

Hierarquia (uma fonte por artigo):
  1. wos_correspondencia — coluna `affiliation` (exportacao da Web of Science): a primeira linha marcada
     «(Corresponding Author)»; o pais e o ultimo elemento dessa afiliacao.
  2. bdtd — teses e dissertacoes da BDTD: Brasil.
  3. crossref_primeiro_autor — sem afiliacao no banco: primeira afiliacao do primeiro autor com afiliacao no
     registo Crossref do DOI (o Crossref nao marca o autor de correspondencia).
  4. pdf — os restantes: leitura da p.1 do PDF (lotes_paises/resultado_loteNN.json), autor de correspondencia;
     sem marca de correspondencia, o primeiro autor; em teses, a instituicao do autor.

Escreve em articles_final.xlsx (aba Articles): col 57 pais_correspondencia, 58 pais_iso3, 59 pais_fonte,
60 pais_citacao (afiliacao usada, tal como impressa).
"""
import glob
import json
import re
import sys

import openpyxl

sys.path.insert(0, 'lotes_paises')
import paises  # noqa: E402

BASE = 'articles_final.xlsx'
COLS = {57: 'pais_correspondencia', 58: 'pais_iso3', 59: 'pais_fonte', 60: 'pais_citacao'}
# afiliacao do Crossref sem pais escrito: pais da instituicao nomeada
CROSSREF_INSTITUICAO = {
    2785: ('Spain', 'Universitat Jaume I (Castello, Espanha)'),
    4930: ('South Korea', 'Seoul National University'),
    4931: ('USA', 'Oregon State University'),
    4993: ('Germany', 'Technical University of Munich'),
}


# Mais de um autor de correspondencia no PDF: usa-se o primeiro na ordem dos autores (mesma regra da
# coluna affiliation da WoS, onde se usa a primeira linha «Corresponding Author»).
PDF_PRIMEIRO_NA_ORDEM = {
    3170: ('India', 'N. Anand (** no PDF; 3.o autor, antes de E. Lubloy): Karunya Institute of Technology and '
                    'Sciences, Coimbatore, India (p.1)'),
}


def clean(s):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(s))


def wos(aff):
    segs = [m.group(1) for m in re.finditer(
        r'\(Corresponding Author\),\s*(.*?)\.(?=\s*[A-Z][^,]*,|\s*$|\n)', str(aff), re.S)]
    if not segs:
        return None
    last = segs[0].split(',')[-1].strip()
    if re.search(r'\bUSA$', last):
        last = 'USA'
    return last, segs[0]


def crossref(w, row):
    au = [x for x in (w or {}).get('author', []) if x.get('affiliation')]
    if not au:
        return None
    s = au[0]['affiliation'][0]['name']
    if row in CROSSREF_INSTITUICAO:
        return CROSSREF_INSTITUICAO[row][0], s + ' [pais da instituicao]'
    for t in reversed([t.strip() for t in re.split(r',|\s{2,}', s)]):
        if paises.normaliza(t):
            return t, s
    for k in paises.PAISES:
        if re.search(r'\b' + re.escape(k) + r'\b', s):
            return k, s
    return None


def main():
    arts = json.load(open('lotes_paises/arts.json'))
    cr = json.load(open('lotes_paises/crossref.json'))
    pdf = {}
    for f in sorted(glob.glob('lotes_paises/resultado_lote*.json')):
        for r in json.load(open(f)):
            pdf[r['row_id']] = r
    wb = openpyxl.load_workbook(BASE)
    ws = wb['Articles']
    for c, nome in COLS.items():
        assert ws.cell(1, c).value in (None, nome), (c, ws.cell(1, c).value)
        ws.cell(1, c).value = nome
    falta, cont = [], {}
    for a in arts:
        r = a['row']
        assert ws.cell(r, 20).value == 'Elegível', r
        res = None
        if a['aff']:
            x = wos(a['aff'])
            if x:
                res = (x[0], 'wos_correspondencia', x[1])
        elif a['source'] == 'BDTD':
            res = ('Brazil', 'bdtd', 'tese/dissertacao BDTD')
        else:
            x = crossref(cr.get(a['doi']), r)
            if x:
                res = (x[0], 'crossref_primeiro_autor', x[1])
            elif r in PDF_PRIMEIRO_NA_ORDEM:
                res = (PDF_PRIMEIRO_NA_ORDEM[r][0], 'pdf_correspondencia', PDF_PRIMEIRO_NA_ORDEM[r][1])
            elif r in pdf and pdf[r].get('pdf_confere') and pdf[r].get('pais'):
                p = pdf[r]
                res = (p['pais'], 'pdf_' + p.get('criterio', ''),
                       f"{p.get('autor', '')}: {p.get('afiliacao', '')} ({p.get('pagina', '')})")
        norm = paises.normaliza(res[0]) if res else None
        if not norm:
            falta.append((r, res))
            continue
        ws.cell(r, 57).value = norm[0]
        ws.cell(r, 58).value = norm[1]
        ws.cell(r, 59).value = res[1]
        ws.cell(r, 60).value = clean(res[2])[:500]
        cont[res[1]] = cont.get(res[1], 0) + 1
    wb.save(BASE)
    print(cont, 'total', sum(cont.values()))
    print('sem pais:', falta)


if __name__ == '__main__':
    main()

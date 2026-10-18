"""Aplica a fronteira do sistema e o transporte (F37, Kie 2026-10-18) as linhas de dados_fck (meta_analysis_base.xlsx).

Colunas novas:
  60 fronteira             A1-A3 | A1-A4 | A1-A5 | A1-A3+C+D | outra | nao_declarada | nao_lido
  61 fronteira_texto       termo do artigo (cradle-to-gate, cradle-to-grave...)
  62 transporte_materiais  incluido | excluido | parcial | nao_declarado | nao_lido   (A2: materias-primas -> central)
  63 transporte_betao_a4   incluido | excluido | nao_declarado | nao_lido            (central -> obra)
  64 transporte_dados      distancias, modal, camiao, factor (texto)
  65 base_processos        processos/bases citados
  66 fronteira_citacao     pagina/tabela/figura + frase
  67 fronteira_nota        notas do leitor (inclui o aviso de «figura nao vista»)
  68 leitura_fronteira     texto+figuras | so_texto | nao_lido  (derivado da nota)
  69 a4_no_valor           sim | nao | incerto  (por confirmar: derivado de col 63; ver DECISOES F37)
Por par: `por_par` do lote sobrepoe o estudo (campos iguais). Uso: python3 lotes_fronteira/aplicar.py [--seco]
"""
import glob
import json
import re
import sys

import openpyxl

BASE = 'meta_analysis_base.xlsx'
NOVAS = {60: 'fronteira', 61: 'fronteira_texto', 62: 'transporte_materiais', 63: 'transporte_betao_a4',
         64: 'transporte_dados', 65: 'base_processos', 66: 'fronteira_citacao', 67: 'fronteira_nota',
         68: 'leitura_fronteira', 69: 'a4_no_valor'}


def clean(s, n=1500):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(s))[:n] if s not in (None, '') else None


def leitura(r):
    nota = (r.get('nota') or '').lower()
    if r.get('texto_completo_lido') is False or 'nao_lido' in nota[:40]:
        return 'nao_lido'
    sem_fig = ('nao vi' in nota or 'não vi' in nota or 'sem ver' in nota or 'figuras nao vistas' in nota
               or 'figuras não vistas' in nota or 'read_file_content' in nota or 'so texto' in nota or 'só texto' in nota
               or 'nao renderiz' in nota or 'não renderiz' in nota)
    return 'so_texto' if sem_fig else 'texto+figuras'


def dados(d):
    if not isinstance(d, dict):
        return clean(d)
    partes = []
    for k, rot in (('dist_materiais_km', 'materiais km'), ('dist_betao_km', 'betao km'), ('modal', 'modal'),
                   ('camiao', 'camiao'), ('factor', 'factor')):
        if d.get(k) not in (None, ''):
            partes.append(f'{rot}: {d[k]}')
    return '; '.join(partes) or None


def main():
    seco = '--seco' in sys.argv
    res = {}
    for f in sorted(glob.glob('lotes_fronteira/resultado_lote*.json')):
        for r in json.load(open(f, encoding='utf-8')):
            res[int(r['row_id'])] = r
    wb = openpyxl.load_workbook(BASE)
    ws = wb['dados_fck']
    for c, n in NOVAS.items():
        ws.cell(1, c).value = n
    pares = {}
    for r in range(2, ws.max_row + 1):
        rid = ws.cell(r, 1).value
        if rid in res:
            pares.setdefault(rid, []).append(r)
    n_est = n_par = 0
    for rid, r in res.items():
        por = {int(p['sheet_row']): p for p in (r.get('por_par') or []) if p.get('sheet_row')}
        for row in pares.get(rid, []):
            p = {**r, **{k: v for k, v in (por.get(row) or {}).items() if v not in (None, '')}}
            a4 = p.get('transporte_concreto_a4')
            val = {60: p.get('fronteira'), 61: clean(p.get('fronteira_texto'), 300), 62: p.get('transporte_materiais_a1a3'),
                   63: a4, 64: clean(dados(p.get('transporte_dados')), 600), 65: clean(p.get('base_dados_processos'), 800),
                   66: clean(p.get('citacao'), 900), 67: clean(f"{r.get('transporte_materiais_nota') or ''} | {r.get('nota') or ''}", 1500),
                   68: leitura(r), 69: {'incluido': 'sim', 'excluido': 'nao'}.get(a4, 'incerto')}
            for c, v in val.items():
                ws.cell(row, c).value = v
            n_par += 1
        n_est += 1
    print(f'{n_est} estudos, {n_par} linhas de dados_fck')
    if not seco:
        wb.save(BASE)
        print('gravado', BASE)


if __name__ == '__main__':
    main()

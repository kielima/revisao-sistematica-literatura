"""Aplica a releitura do tipo de cimento (F12, decisao da Kie 2026-10-10).

Kie: «todos esses [cimentos] que ja existem em mercado deixe como opc. E dentro da familia do cimento
Portland crie subfamilias (CP I, CP II-E, CEM II, V, PLC, etc.) [...] com um botao de informacao (i)».

Regra (F12): ligante = cimento comercial tal como vendido -> familia `opc`, com subfamilia = tipo de cimento
(lotes_cimento/subfamilia.py); adicao juntada pelo proprio estudo ao traco -> `scm`; artigo que nao diz o
cimento (so factor generico de base de dados) -> fica na familia actual; em `opc` recebe a subfamilia
`nao_declarado`.

Escreve em meta_analysis_base.xlsx, aba dados_fck:
  col 48 cimento_subfamilia | col 49 cimento_designacao (como o artigo escreve) | col 50 cimento_norma
  (a citada pelo artigo) | col 51 cimento_citacao (pagina/tabela/frase + nota do lote)
e as mudancas de familia em lotes_familia/familias_revarredura.json (chave row_id|material).
Correr depois: classificar_base_portland.py, metanalise/preparar_dados.py, exportacao.
"""
import json
import re
import sys

import openpyxl

sys.path.insert(0, 'lotes_cimento')
import subfamilia  # noqa: E402

FAM = 'lotes_familia/familias_revarredura.json'
COLS = {48: 'cimento_subfamilia', 49: 'cimento_designacao', 50: 'cimento_norma', 51: 'cimento_citacao'}


def clean(s):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(s))


def main():
    res = subfamilia.carrega()
    sys.path.insert(0, '.')
    import importlib.util
    spec = importlib.util.spec_from_file_location('mexp', 'metanalise-export.py')
    mexp = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mexp)
    fam = json.load(open(FAM, encoding='utf-8'))
    wb = openpyxl.load_workbook('meta_analysis_base.xlsx')
    ws = wb['dados_fck']
    for c, nome in COLS.items():
        assert ws.cell(1, c).value in (None, nome), (c, ws.cell(1, c).value)
        ws.cell(1, c).value = nome
    mud = {'opc->scm': 0, 'scm->opc': 0}
    for r, v in sorted(res.items()):
        assert str(ws.cell(r, 1).value) == str(v['row_id']), r
        material, titulo = ws.cell(r, 7).value, ws.cell(r, 3).value
        actual = mexp.classify(v['row_id'], material, titulo)
        k = f"{v['row_id']}|{material}"
        nova = actual
        if actual == 'opc' and v['ligante'] == 'substituicao_no_traco':
            nova = 'scm'
        elif actual == 'scm' and v['ligante'] == 'cimento_comercial':
            nova = 'opc'
        if nova != actual:
            fam[k] = nova
            mud[f'{actual}->{nova}'] += 1
            obs = ws.cell(r, 13).value or ''
            ws.cell(r, 13).value = clean(f'{obs} [familia {actual} -> {nova} 2026-10-11 (F12, releitura do cimento): '
                                         f'{v["ligante"]}; {v["designacao"]}]').strip()
        ws.cell(r, 48).value = v['subfamilia'] if nova == 'opc' else None
        ws.cell(r, 49).value = clean(v['designacao']) if v['designacao'] else None
        ws.cell(r, 50).value = clean(v['norma']) if v['norma'] else None
        ws.cell(r, 51).value = clean(f"{v['citacao'] or ''} | nota: {v['nota'] or ''}")[:1500]
    wb.save('meta_analysis_base.xlsx')
    json.dump(fam, open(FAM, 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
    print(len(res), 'pares;', mud)


if __name__ == '__main__':
    main()

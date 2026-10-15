"""F18 — CO2 que nao se consegue levar ao material sai em E2_co2 (decisao da Kie, 2026-10-15).

Kie (2026-10-15): «se e por viga ou por laje, se tiver as dimensoes do elemento da pra fazer a conta e colocar
como concreto armado ou alguma familia. Ja o edificio todo, se ele nao separa, nao da pra converter. Eu nao quero
tirar pilares, vigas, lajes e elementos que da pra converter. Quero tirar so o que nao da pra converter.»
Lista A1 + A2 aprovada pela Kie («Pode», 2026-10-15).

- A1 (34): o CO2 so existe para o edificio, sistema, pavimento/infra-estrutura ou elemento, sem a quantidade de
  betao (volume ou dimensoes) impressa no artigo -> nao converte para kg CO2/m3 de betao.
- A2 (14): o CO2 do artigo nao e do betao/cimento (so uma parcela: brita, activadores, agregado; ou betao fora da
  fronteira) ou nao esta em kg CO2 (so pontuacao ReCiPe endpoint, indices ou percentagens).
- Ficam Elegiveis: CO2 do proprio material por kg/tonelada sem massa volumica (emissao do cimento/betao, como
  pedido) e artigos que falham por outra razao (fck em falta, valores incoerentes).
- Elemento com dimensoes impressas converte (familia concreto_armado): 2666 SAHMENKO (Tab.2: parede de betao
  normal 0,12 m + malha de aco + la mineral, 95,5 kg CO2-eq/m2 berco-ao-portao; fck 35 MPa da Tab.1) = 795,8.
Motivo de cada artigo: releitura integral de 2026-10-11 (lotes_sem_par_revisao/resultado_loteNN.json) e, para
4617/4739, a releitura dos elementos de 2026-10-15 (lotes_elementos/resultado_lote02/03.json).
"""
import glob
import json
import re

import openpyxl

ART = 'articles_final.xlsx'
BASE = 'meta_analysis_base.xlsx'
FAM = 'lotes_familia/familias_revarredura.json'
HOJE = '2026-10-15'

A1 = [2886, 3111, 3328, 4010, 4264, 4923, 5017, 2976, 4369, 4352, 4322, 4172, 4337, 4024, 3488, 3667, 2871, 4572,
      2668, 3208, 3270, 3690, 3560, 3886, 4037, 4664, 2844, 2761, 4102, 4461, 2797, 3280, 4617, 4739]
A2 = [2586, 2780, 2927, 3480, 4387, 4265, 4247, 3594, 4397, 3106, 4603, 2741, 3420, 3135]


def clean(s):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(s))


def motivos():
    m = {}
    for f in sorted(glob.glob('lotes_sem_par_revisao/resultado_lote*.json')):
        for r in json.load(open(f, encoding='utf-8')):
            m[r['row_id']] = r.get('nao_convertivel_motivo') or r.get('notas') or ''
    for f in sorted(glob.glob('lotes_elementos/resultado_lote*.json')):
        for r in json.load(open(f, encoding='utf-8')):
            m[r['row_id']] = r.get('motivo_se_nao') or r.get('notas') or m.get(r['row_id'], '')
    return m


def main():
    mot = motivos()
    wa = openpyxl.load_workbook(ART)
    ws = wa['Articles']
    for grupo, ids, txt in (
            ('A1', A1, 'CO2 so do edificio/sistema/elemento sem a quantidade de betao impressa - nao converte'),
            ('A2', A2, 'CO2 que nao e do betao/cimento ou nao esta em kg CO2 - nao converte')):
        for r in ids:
            assert ws.cell(r, 20).value == 'Elegível' and ws.cell(r, 40).value == 'Aceito', r
            ws.cell(r, 20).value = 'Rejeitado'
            ws.cell(r, 40).value = 'Rejeitado'
            ws.cell(r, 43).value = 2
            ws.cell(r, 44).value = 'E2_co2'
            m = re.sub(r'\s+', ' ', mot.get(r, ''))[:500]
            ws.cell(r, 32).value = clean(f'{ws.cell(r, 32).value or ""} [F18 {HOJE} (Kie: «tirar so o que nao da '
                                         f'pra converter»), grupo {grupo}: E2_co2 Rejeitado - {txt}. {m}]'.strip())
    wa.save(ART)

    wb = openpyxl.load_workbook(BASE)
    wd = wb['dados_fck']
    fam = json.load(open(FAM, encoding='utf-8'))
    r = 143
    assert wd.cell(r, 1).value == 2666 and wd.cell(r, 14).value == 95.5
    wd.cell(r, 34).value = '95,5 kg CO2-eq/m2 de parede (Amostra 5: betao normal 0,12 m + malha de aco + la mineral)'
    wd.cell(r, 35).value = 'Tab.2 p.926 (GWP, espessuras e massas por m2); Tab.1 p.921 (fck 35 MPa); sec.2 (berco-ao-portao)'
    wd.cell(r, 36).value = ('calculado (F18, elemento com dimensoes impressas): 95,5 kg/m2 / 0,12 m3 de betao por m2 = '
                            '795,8 kg CO2-eq/m3. Verificacao: 0,12 m x 2.400 kg/m3 (Tab.1) = 288 kg/m2 = valor impresso. '
                            'Inclui a malha de aco (18 kg/m2) e a la mineral (0,20 m), nao separaveis -> concreto_armado')
    wd.cell(r, 37).value = 795.8
    wd.cell(r, 23).value = 0.12
    wd.cell(r, 24).value = 'medido_elemento_real'
    wd.cell(r, 25).value = 'Tab.2: camada de betao normal de 0,12 m por 1 m2 de parede (unidade funcional = 1 m2)'
    wd.cell(r, 45).value = 'nenhum'
    wd.cell(r, 47).value = 'sim'
    wd.cell(r, 13).value = clean(f'{wd.cell(r, 13).value or ""} [F18 {HOJE}: convertido por 1 m2 de parede / 0,12 m de '
                                 'betao (elemento com dimensoes impressas, regra da Kie); familia concreto_armado; '
                                 'fck 35 MPa = dado do fabricante (Tab.1), sem idade]'.strip())
    fam[f'2666|{wd.cell(r, 7).value}'] = 'concreto_armado'
    wb.save(BASE)
    with open(FAM, 'w', encoding='utf-8') as f:
        json.dump(fam, f, ensure_ascii=False, indent=0)
    print('rejeitados', len(A1) + len(A2))


if __name__ == '__main__':
    main()

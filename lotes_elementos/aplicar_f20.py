"""F20 — CHOI 4620 e MITOULIS 3691 rejeitados em E2_co2 (decisao da Kie, 2026-10-15: «Pode rejeitar os dois»).

Releitura pessoal dos dois PDFs a 2026-10-15 (texto completo pelo conector do Drive):

- 3691 MITOULIS 2023 (Transp. Res. D 121, 103800): o betao so entra como factores ICE de Gibbons & Orr 2022
  (Tab.A1 p.15: C25/30 CEM I 0,142 e C25/30 25% GGBS 0,130 kg CO2e/kg). Os resultados (Tab.A3 p.18) sao tCO2e por
  tarefa de reparacao de uma ponte, somando betao, aco, cofragem, pedra e gabioes, sem o volume de betao de cada
  tarefa; o unico total e 3.516 tCO2e = 2,6 t/m2 de tabuleiro. CO2 de sistema sem volume de betao impresso = F18 A1.
- 4620 CHOI 2020 (Sustainability 12, 9193): o unico CO2 do betao e o factor generico ICE v2.0 0,141 kg CO2-eq/kg
  (p.4), multiplicado pela massa dos paineis. O artigo nao imprime massa volumica nem volume de betao; os 24 kN/m3
  da Tab.1 (p.6) sao o peso proprio usado na carga permanente da laje, nao a massa dos paineis, e a geometria dos
  paineis da 1.577 kg/m3 (incoerente). Converter exigiria assumir a massa volumica. Precedente JATIVA 3069 (F17):
  factor generico de literatura sem ACV propria do betao = Rejeitado, mesmo usado no calculo.
"""
import re

import openpyxl

ART = 'articles_final.xlsx'
HOJE = '2026-10-15'

MOTIVO = {
    3691: ('F18 A1 + factor generico: CO2 so por tarefa de reparacao da ponte (Tab.A3 p.18, betao+aco+cofragem+pedra+'
           'gabioes), sem volume de betao impresso; o betao so entra como factores ICE de Gibbons & Orr 2022 (Tab.A1 '
           'p.15, C25/30 0,142/0,130 kg CO2e/kg); total 3.516 tCO2e = 2,6 t/m2 de tabuleiro - nao converte para kg/m3'),
    4620: ('precedente JATIVA 3069 (F17): o unico CO2 do betao e o factor generico ICE v2.0 0,141 kg CO2-eq/kg (p.4), '
           'sem ACV propria do betao; sem massa volumica nem volume impressos (os 24 kN/m3 da Tab.1 p.6 sao o peso '
           'proprio da carga permanente da laje; a geometria dos paineis da 1.577 kg/m3, incoerente) - converter '
           'exigiria assumir a massa volumica'),
}


def clean(s):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(s))


def main():
    wa = openpyxl.load_workbook(ART)
    ws = wa['Articles']
    for r, m in MOTIVO.items():
        assert ws.cell(r, 20).value == 'Elegível' and ws.cell(r, 40).value == 'Aceito', r
        ws.cell(r, 20).value = 'Rejeitado'
        ws.cell(r, 40).value = 'Rejeitado'
        ws.cell(r, 43).value = 2
        ws.cell(r, 44).value = 'E2_co2'
        ws.cell(r, 32).value = clean(f'{ws.cell(r, 32).value or ""} [F20 {HOJE} (Kie: «Pode rejeitar os dois»): '
                                     f'E2_co2 Rejeitado - {m}]'.strip())
    wa.save(ART)
    print('rejeitados', len(MOTIVO))


if __name__ == '__main__':
    main()

"""F18 — 5046 SILVA 2022 (dissertacao UFRGS), verificada pessoalmente a 2026-10-15 depois da aprovacao da regra.

Tab.16 (p.166): GWP do sistema «Estruturas» = 5,76E+04 kg CO2 eq (A1-A3, A4, B4, C2). Tab.11 (pp.143-144): o sistema
junta 274,60 m3 de betao usinado fck 25 MPa com 16.143,99 kg de aco, cofragem (38,04 m3 de madeira + compensado),
lajotas ceramicas, EPS e um betao/argamassa feito em obra (750 kg de CP II-32, 3.675 kg de areia, 7.840 kg de brita)
cujo volume nao e impresso. O CO2 do processo de betao nunca e dado em kg (so % por sistema). Sem o volume total de
betao do sistema, total / 274,60 m3 nao e o CO2 por m3 de betao (alem disso daria ~210 kg/m3 com o aco incluido,
abaixo do que o proprio inventario implica). Grupo A1 da regra F18: nao converte -> Rejeitado em E2_co2.
"""
import re

import openpyxl

ART = 'articles_final.xlsx'
r = 5046
wa = openpyxl.load_workbook(ART)
ws = wa['Articles']
assert ws.cell(r, 20).value == 'Elegível' and ws.cell(r, 40).value == 'Aceito'
ws.cell(r, 20).value = 'Rejeitado'
ws.cell(r, 40).value = 'Rejeitado'
ws.cell(r, 43).value = 2
ws.cell(r, 44).value = 'E2_co2'
txt = ('[F18 2026-10-15 (Kie: «tirar so o que nao da pra converter»), grupo A1, verificado pessoalmente: E2_co2 '
       'Rejeitado - GWP so do sistema Estruturas (Tab.16, 5,76E+04 kg CO2 eq, A1-A3/A4/B4/C2), que junta 274,60 m3 de '
       'betao usinado com aco, cofragem, lajotas ceramicas, EPS e betao de obra sem volume impresso (Tab.11); CO2 do '
       'processo de betao nunca em kg; nao converte]')
ws.cell(r, 32).value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', f'{ws.cell(r, 32).value or ""} {txt}'.strip())
wa.save(ART)
print('ok')

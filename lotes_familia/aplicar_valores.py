"""Correccoes de valor achadas na revarredura das familias e do betao armado (2026-10-04).

Cada entrada foi confirmada no PDF (texto extraido e/ou figura) — a coluna `verif` diz como:
  'pessoal'   = conferido por mim no texto/figura do artigo;
  'subagente' = leitura do subagente da revarredura com citacao exacta, nao reconferida por mim
                (ficam marcadas «por confirmar» em metanalise/DECISOES_METANALISE.md).
Regras aplicadas (DECISOES_METANALISE.md, seccao D):
  D1  CO2 com aco separavel -> grava-se o CO2 so do betao (e a linha sai de concreto_armado);
  E1  valor que nao e o CO2 do betao (so a brita, so o gasoleo, sem o cimento, sem os activadores,
      soma de factores, por ciclo de durabilidade) -> NAO CONVERTIDO (fora do grafico);
  E2  linha que repete o mesmo traco/par de outra linha do mesmo artigo -> NAO CONVERTIDO (duplicado);
  E3  unidade por m3 normalizado a resistencia gravada como kg/m3 -> x fck (regra total+divisor).
Correr da raiz do repo; depois `lotes_familia/aplicar.py`, `classificar_base_portland.py` e `recalcular_ci.py`.
"""
import re

import openpyxl

HOJE = '2026-10-04'
BASE = 'meta_analysis_base.xlsx'

# linha: (row_id, co2_velho, co2_novo, co2_original, local, conversao, verif, nota)
CO2 = {
    135: (2684, 292.8, 311.6, '~311,6 kg CO2-eq/m3 de betao (Fig.5)', 'Fig.5 p.6', 'nenhuma - lido por pixel na Fig.5 (por m3 de betao)', 'pessoal',
          'D1: o 292,8 era 58,56 kg/m2 de laje armada (Tab.5, inclui varao) / 0,20 m3; a Fig.5 da o betao por m3 (barra OPC)'),
    136: (2684, 357.2, 386.6, '~386,6 kg CO2-eq/m3 de betao (Fig.5)', 'Fig.5 p.6', 'nenhuma - lido por pixel na Fig.5 (por m3 de betao)', 'pessoal',
          'D1: o 357,2 era 71,44 kg/m2 de laje (Tab.5, inclui varao) / 0,20 m3; Fig.5 barra OPC-SF (fibras dispersas, nao armadura)'),
    137: (2684, 186.23, 218.8, '~218,8 kg CO2-eq/m3 de betao (Fig.5)', 'Fig.5 p.6', 'nenhuma - lido por pixel na Fig.5 (por m3 de betao)', 'pessoal',
          'D1: o 186,23 era 40,97 kg/m2 de laje (Tab.5, inclui varao) / 0,22 m3; Fig.5 barra LC3'),
    138: (2684, 249.82, 290.7, '~290,7 kg CO2-eq/m3 de betao (Fig.5)', 'Fig.5 p.6', 'nenhuma - lido por pixel na Fig.5 (por m3 de betao)', 'pessoal',
          'D1: o 249,82 era 54,96 kg/m2 de laje (Tab.5, inclui varao) / 0,22 m3; Fig.5 barra LC3-SF. Nota: a Tab.5 descontado o varao nao reproduz a Fig.5 (incoerencia interna do artigo); usa-se a Fig.5, que e por m3 de betao'),
    646: (2693, 304.2, 285.0, '14,6 kg CO2-eq por viga x 93,7% (betao)', 'Tab.3 + texto Sec.3', '14,6 x 0,937 / 0,048 m3', 'subagente',
          'D1: o 304,2 inclui o varao BFRP; o texto da a parcela do betao (93,7%)'),
    647: (2693, 303.8, 284.6, 'total da viga x 93,7% (betao)', 'Tab.3 + texto Sec.3', 'total x 0,937 / volume da viga', 'subagente', 'D1: idem 646'),
    648: (2693, 303.7, 284.6, 'total da viga x 93,7% (betao)', 'Tab.3 + texto Sec.3', 'total x 0,937 / volume da viga', 'subagente', 'D1: idem 646'),
    649: (2693, 303.8, 284.6, 'total da viga x 93,7% (betao)', 'Tab.3 + texto Sec.3', 'total x 0,937 / volume da viga', 'subagente', 'D1: idem 646'),
    650: (2693, 501.6, 281.9, '23,7 kg CO2-eq por viga x 56,2% (betao)', 'Tab.3 + texto Sec.3', '23,7 x 0,562 / 0,04725 m3', 'subagente',
          'D1: o 501,6 inclui 4,44 kg de varao de aco (43,8%)'),
    234: (2801, 261.7, 239.4, '22,5 kg CO2-eq (concreto, 70,09%)', 'Tabela 15 p.88', '22,5 kg / 0,094 m3 de concreto C20', 'pessoal',
          'D1: o 261,7 usava um factor 0,771 sem fonte no PDF e ficava entre o concreto (22,5) e concreto+vergalhao (25,31)'),
    331: (3276, 613.0, 265.22, '265,22 kg CO2-eq/m3 (Concrete, normal, at plant)', 'Table 3', 'nenhuma - factor por m3 de betao do inventario do artigo', 'pessoal',
          'D1: o 613 era 153,24 kg/m2 de AREA DE PISO (inclui aco, cofragem, obra, fim de vida) / 0,25 m; o betao tem factor proprio na Table 3'),
    1150: (3958, 408.0, 253.44, '408,44 - 155,00 (aco) kg CO2e/m3', 'Table 1', '408,44 - 155,00 = 253,44 (inclui pre-fabrico 34,38)', 'pessoal',
           'D1: o 408 inclui 100 kg de armadura (155 kg CO2e) da Table 1'),
    821: (3053, 450.0, 334.9, '(17.222+88+241+2+195) kg CO2-eq / 53 m3', 'Tab.4', 'soma da producao dos constituintes do betao / 53 m3', 'pessoal',
          'D1: o 450 (Fig.5b) inclui a malha de armadura (6.700 kg CO2-eq na Tab.4); so constituintes do betao, sem transporte'),
    305: (3073, 295.0, 215.4, '6,0 kg CO2/MPa (Mix-1, Fig.2B) x 35,9 MPa', 'Fig.2(B) + Tab. de resistencias', '6,0 x 35,9 = 215,4', 'subagente',
          'par corrigido: o 295 era da Wall-3 (Mix-3, 25,2 MPa, parede com armadura, Fig.4); o fck 35,9 e da Mix-1'),
    312: (3092, 500.0, 498.87, '498,87 CO2-eq', 'texto Sec. resultados', 'nenhuma - publicado em kg CO2-eq/m3', 'pessoal', 'arredondamento corrigido'),
    298: (3071, 222.0, 108.0, '2,7 kg CO2/MPa (Table 4) x 40 MPa', 'Table 4', '2,7 x 40 = 108', 'subagente',
          'o 222 da Table 4 contradiz o proprio artigo (Table A4 repete a Table A1 de energia; -73% face ao OPCC; silicato = 54%); 108 = indicador por MPa impresso x fck'),
    300: (3071, 512.0, 259.0, '0,259 t-CO2/m3 (soma da Table A5)', 'Table A5', '0,259 t = 259 kg', 'subagente',
          'o 512 da Table 4 contradiz a Table A5, o indicador por MPa (4,4 x 58,5 = 257) e o texto (OPCC 68% maior)'),
    303: (3071, 639.0, 313.0, '0,313 t-CO2/m3 (soma da Table A6)', 'Table A6', '0,313 t = 313 kg', 'subagente',
          'o 639 da Table 4 contradiz a Table A6, o indicador por MPa (3,1 x 102 = 316) e o texto'),
    1222: (4201, 496.0, 545.6, '496 kg CO2-eq/m3-eq (M2), FU = 1,10 m3-eq', 'texto + Table 2', '496 x 1,10 = 545,6 (por m3 real; so A1)', 'subagente',
           'a unidade do artigo e por m3-EQUIVALENTE (FU ajustada pela resistencia, Eq.1-2)'),
    1223: (4201, 383.0, 459.6, '383 kg CO2-eq/m3-eq (M3), FU = 1,20 m3-eq', 'texto + Table 2', '383 x 1,20 = 459,6 (por m3 real; so A1)', 'subagente',
           'idem 1222'),
    1494: (4521, 66.6, 64.0, '64 kg CO2eq/m3 (sc10)', 'Fig.8', 'nenhuma - lido na Fig.8', 'subagente',
           'o 66,6 era o PII arredondado (0,9) x fck; a Fig.8 da o valor directo'),
    976: (3384, 334.4, 350.7, '~350,7 kg CO2-eq/m3 (CC)', 'Fig.3', 'nenhuma - lido por pixel na Fig.3', 'pessoal',
          'o 334,4 vinha da Fig.2 normalizada (8E-12 x 4,18E13, 1 algarismo)'),
    977: (3384, 418.0, 449.7, '~449,7 kg CO2-eq/m3 (SFRCC)', 'Fig.3', 'nenhuma - lido por pixel na Fig.3', 'pessoal', 'idem 976'),
    978: (3384, 167.2, 185.7, '~185,7 kg CO2-eq/m3 (S_1)', 'Fig.3', 'nenhuma - lido por pixel na Fig.3', 'pessoal', 'idem 976'),
    979: (3384, 167.2, 166.9, '~166,9 kg CO2-eq/m3 (GF_5)', 'Fig.3', 'nenhuma - lido por pixel na Fig.3', 'pessoal', 'idem 976'),
    980: (3384, 167.2, 179.7, '~179,7 kg CO2-eq/m3 (PP_4)', 'Fig.3', 'nenhuma - lido por pixel na Fig.3', 'pessoal', 'idem 976'),
    1826: (4869, 337.67, 302.39, '337,67 - 35,28 (aco: producao, preparacao e transportes) kg CO2/m3', 'Tabla 2', '337,67 - 35,28 = 302,39', 'subagente',
           'D1: o 337,67 (Med) e de betao armado HA-25 com 50 kg de aco/m3'),
    1748: (4758, 493.89, 358.03, 'calculado com as Tabelas 1 e 3 (silica ativa 3,4E-3 kg CO2/kg)', 'Tabelas 1 e 3', 'soma traco x factores', 'subagente',
           'o calculo antigo usou 3,4 kg CO2/kg para a silica ativa; a Tabela 3 da 3,4E-3 (+135,86 kg/m3)'),
    1752: (4758, 470.53, 334.67, 'calculado com as Tabelas 1 e 3 (silica ativa 3,4E-3 kg CO2/kg)', 'Tabelas 1 e 3', 'soma traco x factores', 'subagente', 'idem 1748'),
    1756: (4758, 447.18, 311.31, 'calculado com as Tabelas 1 e 3 (silica ativa 3,4E-3 kg CO2/kg)', 'Tabelas 1 e 3', 'soma traco x factores', 'subagente', 'idem 1748'),
    51: (5041, 515.8, 224.6, '50,95 kg CO2-eq/m2 (betao: brita, areia, cimento)', 'Fig.29 p.136', '50,95 kg/m2 / volume de betao por m2 (Quadros 2-7)', 'subagente',
         'D1: o 515,8 e o total da estrutura (Quadro 8) com aco, cofragens, distribuicao e deposicao'),
    43: (5018, 494.7, 198.4, '40,1% x 48.485,13 kg CO2-eq (betao 25 MPa, texto p.96)', 'texto p.96 + total da estrutura', '0,401 x 48.485,13 / 98,0 m3', 'subagente',
         'D1: o 494,7 era o total da estrutura (aco, madeira, transporte) / volume de betao'),
}

# linha: (row_id, fck_velho, fck_novo, nota)
FCK = {
    976: (3384, 38, 35, 'Tab.1 CC = 35 MPa (ref.[1]); o 38 era a faixa "38 +/- 2" do texto'),
    977: (3384, 38, 39.5, 'Tab.1 SFRCC = 39,5 MPa (ref.[1])'),
    978: (3384, 38, 38.4, 'Tab.1 S_1 = 38,4 MPa (ref.[1])'),
    979: (3384, 38, 35.97, 'Tab.1 GF_5 = 35,97 MPa (ref.[6])'),
    980: (3384, 38, 39.21, 'Tab.1 PP_4 = 39,21 MPa (ref.[7])'),
}

# linha: (row_id, motivo) -> NAO CONVERTIDO (sai do grafico)
NAO_CONV = {
    4: (2586, 'E1: 30,6 e so o GWP da brita do betao C30/37 (1.092 kg x 0,028), nao o CO2 do betao; o artigo nao da o CO2 total por m3'),
    90: (2611, 'E1: 19,944 (Table 14) e o CO2 do GPC dividido por 18 ciclos de 50 anos em acido sulfurico; o valor por m3 so existe na Fig.8'),
    177: (2741, 'E1: 32,24 kg CO2/m3 e so o gasoleo da maquinaria de betonagem (Sec.2.3.2), nao o CO2 do betao'),
    793: (2927, 'E1: o CO2 exclui o cimento ("cement production ... excluded from assessment", Tabela 6)'),
    794: (2927, 'E1: o CO2 exclui o cimento ("cement production ... excluded from assessment", Tabela 6)'),
    880: (3170, 'E1: 3,2628 e a soma dos factores unitarios da Tab.4 (t-CO2/t), nao o CO2 da mistura'),
    423: (3480, 'E1: 92,15 (Tab.8) exclui os activadores NaOH/Na2SiO3 ("do not incorporate water and admixtures")'),
    991: (3409, 'E1: tabuleiro com 11,5 t de armadura e escala do artigo internamente incoerente (10^5 vs 10^4); sem CO2 do betao recuperavel'),
    1106: (3818, 'E2/D1: RC com aco (885,13 = 318,13 + 567,01 do varao); o mesmo traco sem aco ja esta na linha 1107 (RC-0)'),
    151: (2703, 'E2: duplica a linha 154 (mesmo traco 100% OPC, mesmo fck e CO2)'),
    152: (2703, 'E2: duplica a linha 158 (mesmo traco OPC+BFS 30%, mesmo fck e CO2)'),
    153: (2703, 'E2: duplica a linha 157 (mesmo traco OPC+W-BFS 30%, mesmo fck e CO2)'),
    1418: (4476, 'E2: o proprio artigo repete a Mix 6 como Mix 28 (mesma composicao, fck e CO2) - fica a linha 1396'),
    1423: (4476, 'E2: o proprio artigo repete a Mix 15 como Mix 33 (mesma composicao, fck e CO2) - fica a linha 1405'),
    1743: (4754, 'E1: 15,8 kg CO2/m3 impossivel para 182,6 kg de cimento/m3; o eixo da Fig.4 nao tem unidade'),
    1744: (4754, 'E1: 20,1 kg CO2/m3 impossivel para 182,6 kg de cimento/m3; o eixo da Fig.4 nao tem unidade'),
    1761: (4766, 'E2: mesma classe 5500 psi e mesmo factor da linha 1758 (dois elementos do mesmo edificio)'),
    63: (5050, 'E2: mesmo factor generico da DAP Votorantim (25 MPa x 267) da linha 62, aplicado a outro sistema'),
    66: (5055, 'E2: mesmo traco e mesmo CO2 da linha 65, com o fck aos 91 d; fica o par aos 28 d'),
    68: (5055, 'E2: mesmo traco e mesmo CO2 da linha 67, com o fck aos 91 d; fica o par aos 28 d'),
    70: (5055, 'E2: mesmo traco e mesmo CO2 da linha 69, com o fck aos 91 d; fica o par aos 28 d'),
    72: (5055, 'E2: mesmo traco e mesmo CO2 da linha 71, com o fck aos 91 d; fica o par aos 28 d'),
    74: (5055, 'E2: mesmo traco e mesmo CO2 da linha 73, com o fck aos 91 d; fica o par aos 28 d'),
    76: (5055, 'E2: mesmo traco e mesmo CO2 da linha 75, com o fck aos 91 d; fica o par aos 28 d'),
    1836: (5029, 'E1: as Tabelas 27-29 dao totais da ciclovia (588 m3) com os valores uma linha acima do rotulo; o numero gravado e a eutrofizacao, nao o GWP por m3'),
    1837: (5029, 'E1: as Tabelas 27-29 dao totais da ciclovia (588 m3) com os valores uma linha acima do rotulo; o numero gravado e a eutrofizacao, nao o GWP por m3'),
    1838: (5029, 'E1: as Tabelas 27-29 dao totais da ciclovia (588 m3) com os valores uma linha acima do rotulo; o numero gravado e a eutrofizacao, nao o GWP por m3'),
    1839: (5029, 'E1: as Tabelas 27-29 dao totais da ciclovia (588 m3) com os valores uma linha acima do rotulo; o numero gravado e a eutrofizacao, nao o GWP por m3'),
    1840: (5029, 'E1: as Tabelas 27-29 dao totais da ciclovia (588 m3) com os valores uma linha acima do rotulo; o numero gravado e a eutrofizacao, nao o GWP por m3'),
    1841: (5029, 'E1: as Tabelas 27-29 dao totais da ciclovia (588 m3) com os valores uma linha acima do rotulo; o numero gravado e a eutrofizacao, nao o GWP por m3'),
}

# fck sem apoio no PDF -> par retirado (fck apagado)
FCK_SEM_APOIO = {
    1478: (4511, 'o fck 30 nao aparece no PDF: o "25-30 MPa" e o nome do processo Ecoinvent da electricidade; a resistencia esta noutro artigo (ref.[16])'),
    1479: (4511, 'idem 1478'),
}

# rotulos de material errados: linha -> (row_id, texto_errado, texto_certo)
ROTULO = {r: (5055, 'cinza hidratada', 'cal hidratada') for r in (71, 72)}

# E3: PUSHKAR 4580 - gravado em kg CO2-eq/(m3.MPa) (FU = 1 m3 normalizado a resistencia aos 28 d) -> x fck
PUSHKAR_4580 = range(1570, 1590)


def clean(s):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(s))


def nota(ws, r, txt):
    v = ws.cell(r, 13).value
    ws.cell(r, 13).value = clean(f'{v} [{txt}]' if v else f'[{txt}]')


def main():
    wb = openpyxl.load_workbook(BASE)
    ws = wb['dados_fck']
    for r, (rid, velho, novo, orig, local, conv, verif, txt) in CO2.items():
        assert str(ws.cell(r, 1).value) == str(rid) and abs(float(ws.cell(r, 37).value) - velho) < 0.01, (r, ws.cell(r, 37).value)
        ws.cell(r, 37).value = novo
        ws.cell(r, 14).value = novo
        ws.cell(r, 15).value = 'kg CO2-eq/m3'
        ws.cell(r, 34).value = clean(orig)
        ws.cell(r, 35).value = clean(local)
        ws.cell(r, 36).value = clean(conv)
        nota(ws, r, f'CO2 {velho} -> {novo} {HOJE} (revarredura familias/betao armado, verificacao {verif}): {txt}')
    for r, (rid, velho, novo, txt) in FCK.items():
        assert str(ws.cell(r, 1).value) == str(rid) and float(ws.cell(r, 8).value) == velho, r
        ws.cell(r, 8).value = novo
        nota(ws, r, f'fck {velho} -> {novo} {HOJE} (revarredura): {txt}')
    for r, (rid, txt) in NAO_CONV.items():
        assert str(ws.cell(r, 1).value) == str(rid), r
        ws.cell(r, 36).value = clean(f'NAO CONVERTIDO - {txt}')
        ws.cell(r, 37).value = None
        nota(ws, r, f'PAR RETIRADO DO GRAFICO {HOJE} (revarredura): {txt}')
    for r, (rid, txt) in FCK_SEM_APOIO.items():
        assert str(ws.cell(r, 1).value) == str(rid), r
        velho = ws.cell(r, 8).value
        ws.cell(r, 8).value = None
        ws.cell(r, 10).value = None
        nota(ws, r, f'fck {velho} apagado {HOJE} (revarredura): {txt}')
    for r, (rid, errado, certo) in ROTULO.items():
        assert str(ws.cell(r, 1).value) == str(rid) and errado in ws.cell(r, 7).value, r
        ws.cell(r, 7).value = ws.cell(r, 7).value.replace(errado, certo)
        nota(ws, r, f'rotulo "{errado}" -> "{certo}" {HOJE} (revarredura: CH = cal hidratada na tese)')
    for r in PUSHKAR_4580:
        assert str(ws.cell(r, 1).value) == '4580', r
        fck, v = float(ws.cell(r, 8).value), float(ws.cell(r, 37).value)
        assert v < 5, (r, v)
        novo = round(v * fck, 2)
        ws.cell(r, 34).value = f'{v} kg CO2-eq por m3 normalizado a resistencia aos 28 d (FU do artigo)'
        ws.cell(r, 36).value = f'{v} x fck {fck} = {novo} [ALERTA ETAPA4: valor derivado - o inventario usa "cement mortar, at plant/CH" como proxy do CEM I]'
        ws.cell(r, 37).value = novo
        ws.cell(r, 14).value = novo
        ws.cell(r, 15).value = 'kg CO2-eq/m3'
        nota(ws, r, f'CO2 {v} -> {novo} {HOJE} (revarredura, E3): o valor gravado era por m3 E por MPa (FU = 1 m3 normalizado a resistencia); '
                    f'o RELABEL de 2026-09-08 estava errado')
    wb.save(BASE)
    print('ok', len(CO2), len(FCK), len(NAO_CONV), len(FCK_SEM_APOIO), len(PUSHKAR_4580))


if __name__ == '__main__':
    main()

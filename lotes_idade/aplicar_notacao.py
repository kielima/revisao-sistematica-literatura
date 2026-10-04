"""Regra da notacao de norma (decisao da Kie, 2026-10-04).

Classe de resistencia escrita no artigo na notacao de uma norma que a define aos 28 d
(EN 206: Cxx/yy, incluindo 'Cast in Situ 25/30'; EHE-08: HA-xx) conta como idade de 28 d,
mesmo que o artigo nao cite a norma. Valores soltos em MPa (ou 'C30' sem barra) continuam vazios.

Aplicado aos pares 'classe_sem_idade' dos lotes de idade cuja notacao foi confirmada na
citacao do subagente. Excluidos 950/951 (3263 SKULLESTAD): o fck gravado e o valor
'assumido' (30) e a media estimada pelo fornecedor (45), nao o fck da classe da Tab.5.
So preenche idade_dias vazia. Correr da raiz do repo.
"""
import openpyxl

BASE = 'meta_analysis_base.xlsx'
PARES = {  # sheet_row: (row_id, notacao impressa)
    4: ('2586', 'C30/37'), 220: ('2766', 'C25/30'), 221: ('2766', 'C25/30'),
    246: ('2946', 'C30/37'), 247: ('2946', 'C30/37'),
    538: ('3882', 'C35/45'), 539: ('3882', 'C90/105'), 540: ('3882', 'C90/105'),
    629: ('2641', 'C25/30'), 630: ('2650', 'C25/30'),
    795: ('2937', '25/30'), 796: ('2937', '25/30'), 797: ('2937', '25/30'),
    821: ('3053', 'C30/37'), 1013: ('3449', 'C30/37'), 2000: ('3449', 'C30/37'),
    1126: ('3872', 'C35/45'), 1127: ('3872', 'C35/45'), 1128: ('3872', 'C35/45'),
    1140: ('3915', 'C25/30'), 1206: ('4144', 'HA-25'),
    1452: ('4493', 'C70/85'), 1453: ('4493', 'C35/45'),
    1623: ('4605', 'HA-30'), 1624: ('4605', 'HA-30'),
    1688: ('4691', 'C25/30'), 1689: ('4691', 'C25/30'), 1690: ('4691', 'C25/30'),
    1855: ('2687', 'C30/37'),
}


def main():
    wb = openpyxl.load_workbook(BASE)
    ws = wb['dados_fck']
    n = 0
    for r, (rid, notacao) in PARES.items():
        assert str(ws.cell(r, 1).value) == rid, (r, rid)
        if ws.cell(r, 10).value not in (None, ''):
            continue
        ws.cell(r, 10).value = 28
        norma = 'EHE-08' if notacao.startswith('HA') else 'EN 206'
        carimbo = (f'[idade 28 d por notacao de norma (regra Kie 2026-10-04): classe {notacao} '
                   f'impressa no artigo, notacao da {norma}, que define a resistencia aos 28 d]')
        obs = ws.cell(r, 13).value
        ws.cell(r, 13).value = f'{obs} {carimbo}' if obs else carimbo
        n += 1
    wb.save(BASE)
    print('gravadas', n)


if __name__ == '__main__':
    main()

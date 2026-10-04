"""Aplica os resultados dos lotes de idade do ensaio a dados_fck (col 10 idade_dias).

2026-10-04. Regras:
  - so preenche idade_dias quando a celula esta vazia (nunca apaga valor gravado);
  - so grava idade quando o subagente a encontrou ('ensaio' ou 'classe_28d_declarada');
    'classe_sem_idade' e 'nao_declarada' ficam vazias, so com a nota;
  - acrescenta a observacoes (col 13) um carimbo com o tipo e a citacao;
  - SOBREPOR / IGNORAR: correcoes da auditoria pessoal aos resultados dos lotes.
Correr da raiz do repo.
"""
import glob
import json
import re

import openpyxl

BASE = 'meta_analysis_base.xlsx'
COL_ROWID, COL_FCK, COL_IDADE, COL_OBS = 1, 8, 10, 13

# sheet_row -> (idade_dias|None, tipo, citacao) que substitui o resultado do lote
SOBREPOR = {}
IGNORAR = set()  # sheet_rows cujo resultado nao se aplica


def clean_str(s):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(s))


def sem_acentos(s):
    import unicodedata
    return ''.join(c for c in unicodedata.normalize('NFKD', s) if not unicodedata.combining(c))


# 2672 SPOSITO: desdobramento da linha legada de faixa (144) nas 7 pastas da Sec.4.1.3,
# fck aos 28 d (Fig.25 p.87 = Tabela 13 p.88) e CO2 CC+biogenico por m3 (Tabela 12 p.85).
# A linha 145 (NH-SS1.5, 97,21 / 108) ja estava certa; a 144 passa a NH e criam-se 5.
SPOSITO = [  # (material, fck, dp, co2)
    ('NH', 46.44, 1.71, 96.2),
    ('NH-SS1.0', 75.87, 6.53, 104.0),
    ('NH-SS2.0', 90.38, 1.52, 112.0),
    ('NH-BLA1.0', 77.45, 7.79, 93.5),
    ('NH-BLA1.5', 86.35, 9.24, 93.9),
    ('NH-BLA2.0', 61.00, 12.10, 93.9),
]
SPOSITO_OBS = ('[DESDOBRAMENTO 2026-10-04, leitura pessoal do PDF integral] A linha legada guardava '
               'a faixa das 7 misturas (fck 23,07 = NH aos 7 d; CO2 93,5 = minimo da Tabela 12). '
               'Separada num traco por linha: fck aos 28 d da Fig.25 p.87 (igual a Tabela 13 p.88, '
               'usada pela propria tese na ecoeficiencia), CO2 CC+biogenico por m3 da Tabela 12 p.85.')


def desdobrar_sposito(ws):
    src = 144
    assert str(ws.cell(src, 1).value) == '2672'
    ncol = ws.max_column
    linhas = [src] + [ws.max_row + i for i in range(1, len(SPOSITO))]
    for r, (mix, fck, dp, co2) in zip(linhas, SPOSITO):
        if r != src:
            for c in range(1, ncol + 1):
                ws.cell(r, c).value = ws.cell(src, c).value
        ws.cell(r, 7).value = f'Pasta de escoria alcali-ativada (BFS), activador {mix}, Sec.4.1.3'
        ws.cell(r, 8).value = fck
        ws.cell(r, 10).value = 28
        ws.cell(r, 12).value = f'Fig.25 p.87 e Tabela 13 p.88: {mix} 28d {fck} +/- {dp} MPa'
        ws.cell(r, 13).value = SPOSITO_OBS
        ws.cell(r, 14).value = co2
        ws.cell(r, 15).value = 'kg CO2-eq/m3 (F.U. 1 m3, CC + biogenic, Tabela 12 p.85)'
        ws.cell(r, 16).value = 'Tabela 12 p.85'
        ws.cell(r, 26).value = dp
        ws.cell(r, 27).value = 'desvio_padrao'
        ws.cell(r, 28).value = 6
        ws.cell(r, 29).value = f'Fig.25 p.87: {mix} 28d {fck} +/- {dp} (six samples per identification)'
        ws.cell(r, 34).value = f'{co2} kg CO2-eq/m3 (CC + biogenic, F.U. 1 m3)'
        ws.cell(r, 35).value = 'Tabela 12 p.85'
        ws.cell(r, 36).value = 'nenhuma - publicado em kg CO2-eq/m3'
        ws.cell(r, 37).value = co2
        ws.cell(r, 38).value = None
        ws.cell(r, 39).value = None
    return set(linhas)


# correccoes de valor confirmadas por leitura pessoal (2026-10-04)
FCK_CORRIGIR = {
    # 4608 PALUMBO: a Table 3 imprime 'C32/40' e 'C28/35' (Hanson UK) mas a coluna de resistencia
    # da Table 3/4 traz o valor de CUBO (40/35); as outras EPD trazem o de cilindro. fck = cilindro.
    1634: (32, 'C32/40', 'Table 3 p.9-10: EPD7 "C32/40 CEMI Ready Mix Concrete"; a coluna de resistencia '
                         'traz 40 (cubo), fck de cilindro da classe = 32 [corrigido 2026-10-04]'),
    1635: (28, 'C28/35', 'Table 3 p.9-10: EPD8 "C28/35 CIIIA Ready Mix Concrete"; a coluna de resistencia '
                         'traz 35 (cubo), fck de cilindro da classe = 28 [corrigido 2026-10-04]'),
}
TIPO_CORRIGIR = {
    # 4463 SAMEER C50/60: o fck 50 e o de cilindro 150x300 (o volume gravado ja e o do cilindro)
    1380: 'cilindro 150x300mm',
}


def main():
    res = {}
    for f in sorted(glob.glob('lotes_idade/resultado_lote*.json')):
        for art in json.load(open(f, encoding='utf-8')):
            for p in art['pares']:
                res[int(p['sheet_row'])] = (str(art['row_id']), p)
    wb = openpyxl.load_workbook(BASE)
    ws = wb['dados_fck']
    IGNORAR.update(desdobrar_sposito(ws))
    for r, (fck, classe, nota) in FCK_CORRIGIR.items():
        assert str(ws.cell(r, COL_ROWID).value) == '4608'
        ws.cell(r, COL_FCK).value = fck
        ws.cell(r, 9).value = classe
        ws.cell(r, COL_OBS).value = f'{ws.cell(r, COL_OBS).value or ""} [{nota}]'.strip()
    for r, tipo in TIPO_CORRIGIR.items():
        assert str(ws.cell(r, COL_ROWID).value) == '4463'
        ws.cell(r, 11).value = tipo
    gravadas = vazias = ja = 0
    for r, (rid, p) in sorted(res.items()):
        if r in IGNORAR:
            continue
        idade, tipo, cit = p.get('idade_dias'), p.get('tipo'), p.get('citacao') or ''
        if r in SOBREPOR:
            idade, tipo, cit = SOBREPOR[r]
        assert str(ws.cell(r, COL_ROWID).value) == rid, (r, rid, ws.cell(r, COL_ROWID).value)
        if ws.cell(r, COL_IDADE).value not in (None, ''):
            ja += 1
            continue
        if idade is not None and tipo in ('ensaio', 'classe_28d_declarada'):
            ws.cell(r, COL_IDADE).value = int(idade)
            gravadas += 1
        else:
            vazias += 1
        carimbo = clean_str(sem_acentos(f'[idade 2026-10-04: {tipo}; {cit}]'))[:400]
        obs = ws.cell(r, COL_OBS).value
        ws.cell(r, COL_OBS).value = f'{obs} {carimbo}' if obs else carimbo
    wb.save(BASE)
    print(f'gravadas {gravadas} | sem idade {vazias} | ja preenchidas {ja} | total {len(res)}')


if __name__ == '__main__':
    main()

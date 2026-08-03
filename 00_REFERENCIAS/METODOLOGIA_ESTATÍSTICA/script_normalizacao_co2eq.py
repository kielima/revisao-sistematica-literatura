#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Normalizacao CO2eq -- aplica D1-D8."""
import openpyxl, re, csv
from openpyxl.styles import Alignment

NBR_FCK_MIN = 20
FATOR_CUBO_CYL = 0.8
FATOR_PSI_MPA = 1 / 145.038

SRC = "/sessions/sharp-brave-ride/mnt/DISSERTAÇÃO/07_ANÁLISES DE DADOS/REVISÃO SISTEMÁTICA DE LITERATURA/BANCO DE DADOS DA RSL.xlsx"
LOG_CSV = "/sessions/sharp-brave-ride/mnt/DISSERTAÇÃO/07_ANÁLISES DE DADOS/REVISÃO SISTEMÁTICA DE LITERATURA/LOG_NORMALIZACAO_CO2EQ.csv"


def num(s):
    if s is None: return None
    s = str(s).strip().replace(',', '.')
    try: return float(s)
    except ValueError: return None


def parse_fck(fck_str):
    if not fck_str or str(fck_str).strip().upper().startswith('N/A'):
        return []
    s = str(fck_str); s_lower = s.lower()
    results = []

    for m in re.finditer(r'\bC\s*(\d{2,3})\s*/\s*(\d{2,3})\b', s):
        cyl = int(m.group(1))
        if 8 <= cyl <= 200:
            results.append({'fck': cyl, 'tipo_cp': 'cilindro', 'idade': 28,
                            'label': 'C{}/{}'.format(cyl, m.group(2)), 'nota': 'classe_EN206_1'})

    cube_mode = ('cubo' in s_lower) or (' cube' in s_lower) or ('cubico' in s_lower)

    pat_labeled = re.compile(r'([A-Za-z][\w\-%]{0,30}?)\s*[=:]\s*(\d+(?:[.,]\d+)?(?:[eE][+\-]?\d+)?)\s*(?:MPa|N\s*/\s*mm)', re.IGNORECASE)
    for m in pat_labeled.finditer(s):
        v = num(m.group(2))
        if not v or not (1 < v < 500): continue
        label = m.group(1).strip()
        if label.lower() in ('p','pp','fig','figura','tabela','tab','cm','mm'): continue
        tipo_cp = 'cubo->cilindro' if cube_mode else 'cilindro'
        fck_final = round(v * FATOR_CUBO_CYL, 2) if cube_mode else v
        nota = 'cubo->cilindro fator 0,8' if cube_mode else 'parse_labeled'
        if not any(r['label'] == label for r in results):
            results.append({'fck': fck_final, 'tipo_cp': tipo_cp, 'idade': 28, 'label': label, 'nota': nota})

    if not results or len(results) < 2:
        pat_vl = re.compile(r'(\d+(?:[.,]\d+)?)\s*MPa\s*\(([^)]{1,40}?)\)', re.IGNORECASE)
        for m in pat_vl.finditer(s):
            v = num(m.group(1))
            if not v or not (1 < v < 500): continue
            label = m.group(2).strip()
            if 'cubo' in label.lower() or 'cube' in label.lower() or 'kg' in label.lower(): continue
            tipo_cp = 'cubo->cilindro' if cube_mode else 'cilindro'
            fck_final = round(v * FATOR_CUBO_CYL, 2) if cube_mode else v
            nota = 'cubo->cilindro fator 0,8' if cube_mode else 'parse_value_label'
            if not any(r['label'] == label for r in results):
                results.append({'fck': fck_final, 'tipo_cp': tipo_cp, 'idade': 28, 'label': label, 'nota': nota})

    if not results:
        m = re.search(r'(\d+(?:[.,]\d+)?)\s*MPa', s, re.IGNORECASE)
        if m:
            v = num(m.group(1))
            if v and 1 < v < 500:
                tipo_cp = 'cubo->cilindro' if cube_mode else 'cilindro'
                fck_final = round(v * FATOR_CUBO_CYL, 2) if cube_mode else v
                results.append({'fck': fck_final, 'tipo_cp': tipo_cp, 'idade': 28,
                                'label': 'principal', 'nota': 'parse_simples'})

    for m in re.finditer(r'(\d+(?:[.,]\d+)?)\s*psi', s, re.IGNORECASE):
        v = num(m.group(1))
        if v and v > 100:
            mpa = round(v * FATOR_PSI_MPA, 2)
            results.append({'fck': mpa, 'tipo_cp': 'cilindro', 'idade': 28,
                            'label': '{}_psi'.format(v), 'nota': 'psi->MPa'})

    ages = re.findall(r'(\d{1,3})\s*dias?', s_lower)
    age_apl = None
    for a in ages:
        try:
            age = int(a)
            if age in (1,3,7,14,28,56,90,180,365,730):
                age_apl = age
                if age > 28: break
        except ValueError: continue
    if age_apl:
        for r in results:
            r['idade'] = age_apl
            if age_apl > 28:
                r['nota'] = (r.get('nota') or '') + ' | idade>{}d incerteza 15-25%'.format(age_apl)
    return results


def detectar_unidade(s):
    sn = s.lower().replace('³','3').replace('²','2').replace('₂','2').replace(' ','').replace('-eq','eq').replace('co2-e','co2e')
    sn = sn.replace('·','*').replace('•','*')
    if 'co2eq/(m3*mpa)' in sn or 'co2/(m3*mpa)' in sn or '/m3*mpa' in sn or '/m3mpa' in sn:
        return 'damineli', None
    if 'co2eq/mpa' in sn and 'm3' in sn:
        return 'damineli', None
    # tCO2/m3 -- toneladas por m3 (converter para kg multiplicando por 1000)
    if re.search(r'\btco2(?:eq?)?/m3\b', sn):
        return 'por_m3_t', None
    if re.search(r'co2(?:e|eq)?/m3\b', sn):
        return 'por_m3', None
    if re.search(r'co2(?:e|eq)?/m2', sn):
        return 'por_m2', 'UF /m2 -- converter para m3 exige espessura/composicao nao declarada; revisar manualmente'
    if re.search(r'co2(?:e|eq)?/kg', sn):
        return 'por_kg', 'UF /kg -- emissao por massa; nao isolavel a m3 de concreto sem densidade'
    if re.search(r'co2(?:e|eq)?/km', sn):
        return 'por_km', 'UF /km -- exigiria espessura, largura e composicao da camada'
    if re.search(r'co2(?:e|eq)?/fu\b', sn):
        return 'por_fu', 'UF /FU (peca/ponte/edificio); requer extracao caso-a-caso'
    if re.search(r'co2(?:e|eq)?/m\b', sn) and '/m3' not in sn and '/m2' not in sn:
        return 'por_m', 'UF /m linear -- exigiria area de secao transversal'
    if re.search(r'\btco2', sn):
        return 'tonelada', 'tCO2 por peca; UF composita nao convertivel automaticamente'
    return 'desconhecido', None


def parse_emissoes(em_str):
    if not em_str: return []
    s = str(em_str).replace('\xa0', ' ')
    if s.lower().strip().startswith('n/a'): return []
    unidade, flag = detectar_unidade(s)
    pares = []
    pat = re.compile(r'([A-Za-z][\w\-%]{0,30}?)\s*[=:]\s*~?\s*(\d+(?:[.,]\d+)?(?:[eE][+\-]?\d+)?)', re.IGNORECASE)
    for m in pat.finditer(s):
        label = m.group(1).strip()
        val = num(m.group(2))
        if val is None or val <= 0 or val > 1e8: continue
        if label.lower() in ('p','pp','fig','figura','tabela','tab','kg','m','co','cm','mm','reducao','reduction','table','col'): continue
        if re.search(r'\(p\.?\s*' + re.escape(str(int(val))) + r'\b', s, re.IGNORECASE): continue
        pares.append((label, val))
    if not pares:
        for m in re.finditer(r'(\d+(?:[.,]\d+)?)\s*(?:kg\s*)?CO', s, re.IGNORECASE):
            v = num(m.group(1))
            if v and 0 < v < 1e6:
                pares.append(('principal', v)); break

    results = []
    for label, valor in pares:
        if unidade == 'damineli':
            if not (0 < valor < 1e3): continue
            results.append({'valor': valor, 'unidade': 'kgCO2eq/(m3*MPa)',
                            'por_m3': None, 'ci_direto': valor, 'label': label, 'flag_excluir': None})
        elif unidade == 'por_m3':
            if not (0 < valor < 1e4): continue
            results.append({'valor': valor, 'unidade': 'kgCO2eq/m3',
                            'por_m3': valor, 'ci_direto': None, 'label': label, 'flag_excluir': None})
        elif unidade == 'por_m3_t':
            # toneladas: converter para kg (x1000)
            if not (0 < valor < 50): continue  # valores tipicos t/m3: 0.05-2
            kg_val = valor * 1000
            results.append({'valor': valor, 'unidade': 'tCO2eq/m3->kgCO2eq/m3',
                            'por_m3': kg_val, 'ci_direto': None, 'label': label, 'flag_excluir': None})
        elif unidade in ('por_m2','por_kg','por_km','por_fu','por_m','tonelada'):
            results.append({'valor': valor, 'unidade': unidade,
                            'por_m3': None, 'ci_direto': None, 'label': label, 'flag_excluir': flag})
        else:
            results.append({'valor': valor, 'unidade': 'desconhecida',
                            'por_m3': None, 'ci_direto': None, 'label': label,
                            'flag_excluir': 'Unidade nao-extraivel automaticamente; revisar manualmente'})
    seen = set(); deduped = []
    for r in results:
        if r['label'] not in seen:
            deduped.append(r); seen.add(r['label'])
    return deduped


def parse_vida_util(dur_str):
    if not dur_str: return None
    s = str(dur_str).lower()
    if 'n/a' in s and 'vida' not in s: return None
    is_pap = any(k in s for k in ['pap','project analysis period','periodo de analise','horizonte de'])
    is_vu = any(k in s for k in ['vida util','vida útil','design life','service life','durabilidade',
                                  'nbr 15575','rsl =','rsl=','periodo de referencia','reference service life'])
    cand = []
    for m in re.finditer(r'(\d{1,3})\s*anos?', s):
        v = int(m.group(1))
        if 1 <= v <= 500: cand.append(v)
    if not cand: return None
    if is_vu: return max(cand)
    if is_pap and not is_vu: return None
    return max(cand)


# ===== PIPELINE PRINCIPAL =====
def processar():
    wb = openpyxl.load_workbook(SRC)
    ws_an = wb["ANÁLISE"]
    ws_meta = wb["METANÁLISE"]
    log_rows = []
    cnt = {'total': 0, 'col_R': 0, 'col_S': 0, 'sem_ci_com_nota': 0,
           'nao_estrutural': 0, 'uf_incompativel': 0, 'apenas_tracao_flexao': 0, 'meta_linhas': 0}
    atipicos = []

    if ws_meta.max_row > 1:
        for r in range(ws_meta.max_row, 1, -1):
            ws_meta.delete_rows(r)
    meta_row = 2

    for ws_row in range(2, ws_an.max_row + 1):
        estudo = ws_an.cell(row=ws_row, column=1).value
        if not estudo: continue
        cnt['total'] += 1
        material = ws_an.cell(row=ws_row, column=4).value
        durab = ws_an.cell(row=ws_row, column=5).value
        fck_raw = ws_an.cell(row=ws_row, column=6).value
        uf_raw = ws_an.cell(row=ws_row, column=11).value
        em_raw = ws_an.cell(row=ws_row, column=12).value
        fck_list = parse_fck(fck_raw)
        em_list = parse_emissoes(em_raw)
        vu = parse_vida_util(durab)
        coms = []
        ci_R = None; ci_t_S = None; excl = False; mot_excl = None

        fck_na = (not fck_raw) or str(fck_raw).strip().upper().startswith('N/A')
        if fck_na:
            excl = True; mot_excl = 'fck N/A -- nao-concreto-estrutural'
            coms.append('Material sem fck (nao-concreto-estrutural). Mantido para discussao qualitativa.')
            cnt['nao_estrutural'] += 1
        elif fck_list and max(r['fck'] for r in fck_list) < NBR_FCK_MIN:
            excl = True
            fmax = max(r['fck'] for r in fck_list)
            mot_excl = 'fck max ({} MPa) < NBR 6118 ({}MPa)'.format(fmax, NBR_FCK_MIN)
            coms.append('Material nao-estrutural: fck max = {} MPa < NBR 6118 ({} MPa). Mantido para discussao qualitativa.'.format(fmax, NBR_FCK_MIN))
            cnt['nao_estrutural'] += 1

        if not excl and not fck_list:
            sf = str(fck_raw or '').lower()
            if 'flex' in sf or 'trac' in sf or 'tens' in sf:
                excl = True; mot_excl = 'apenas tracao/flexao'
                coms.append('Apenas resistencia a tracao/flexao; sem fck a compressao extraivel.')
                cnt['apenas_tracao_flexao'] += 1

        if not excl and em_list:
            if all(e.get('flag_excluir') for e in em_list):
                excl = True; mot_excl = 'UF incompativel'
                coms.append('UF incompativel: ' + str(em_list[0].get('flag_excluir')) + '. Mantido para discussao qualitativa.')
                cnt['uf_incompativel'] += 1

        if not excl and fck_list and em_list:
            obs = []
            lem = {e['label']: e for e in em_list}
            lfck = {r['label']: r for r in fck_list}
            shared = set(lem) & set(lfck)
            if shared:
                for lab in shared:
                    obs.append((lfck[lab], lem[lab]))
                em_ex = [e for lab,e in lem.items() if lab not in shared]
                fck_ex = [r for lab,r in lfck.items() if lab not in shared]
                if em_ex and len(em_ex) == len(fck_ex):
                    for r,e in zip(fck_ex, em_ex):
                        obs.append((r,e))
            elif len(fck_list) == len(em_list):
                for r,e in zip(fck_list, em_list):
                    obs.append((r,e))
            elif len(fck_list) == 1:
                for e in em_list: obs.append((fck_list[0], e))
            elif len(em_list) == 1:
                for r in fck_list: obs.append((r, em_list[0]))
            else:
                obs.append((fck_list[0], em_list[0]))
                coms.append('AMBIGUIDADE: {} fcks x {} emissoes; primeira observacao usada. Revisar manualmente.'.format(len(fck_list), len(em_list)))
                atipicos.append((str(estudo), '{} fcks x {} emissoes sem casamento'.format(len(fck_list), len(em_list))))

            cis_ok = []
            for fck_obs, em_obs in obs:
                if em_obs.get('flag_excluir'): continue
                if em_obs.get('ci_direto') is not None:
                    ci = em_obs['ci_direto']; conv = 'Damineli direto'
                elif em_obs.get('por_m3') is not None and fck_obs['fck'] > 0:
                    ci = em_obs['por_m3'] / fck_obs['fck']
                    conv = 'emissao/m3 ({}) / fck ({} MPa)'.format(em_obs['por_m3'], fck_obs['fck'])
                else: continue
                # ci_t = ci / vida_util — taxa anual de emissão por unidade de fck
                # kg CO2eq / (m3 . MPa . ano); menor = mais eco-eficiente; premia durabilidade
                ci_t = round(ci / vu, 5) if vu else None
                ci = round(ci, 5)
                cis_ok.append((fck_obs, em_obs, ci, ci_t, conv))
                obs_id = (str(estudo).split(',')[0].strip() + '_' + str(em_obs['label']))[:40].replace(' ','_').replace('.','')
                ws_meta.cell(row=meta_row, column=1, value=obs_id)
                ws_meta.cell(row=meta_row, column=2, value=str(estudo))
                ws_meta.cell(row=meta_row, column=3, value=str(em_obs['label']))
                ws_meta.cell(row=meta_row, column=4, value=fck_obs['fck'])
                ws_meta.cell(row=meta_row, column=5, value=fck_obs['tipo_cp'])
                ws_meta.cell(row=meta_row, column=6, value=fck_obs['idade'])
                ws_meta.cell(row=meta_row, column=7, value=em_obs.get('por_m3') if em_obs.get('por_m3') is not None else '(Damineli)')
                ws_meta.cell(row=meta_row, column=8, value=conv)
                ws_meta.cell(row=meta_row, column=9, value=vu if vu else '')
                ws_meta.cell(row=meta_row, column=10, value=ci)
                ws_meta.cell(row=meta_row, column=11, value=ci_t if ci_t is not None else '')
                ws_meta.cell(row=meta_row, column=12, value='')
                ws_meta.cell(row=meta_row, column=13, value=True)
                notas = []
                if 'cubo' in str(fck_obs.get('nota','')): notas.append('cubo->cilindro fator 0,8')
                if fck_obs.get('idade') and fck_obs['idade'] > 28: notas.append('fck a {}d > 28d incerteza 15-25%'.format(fck_obs['idade']))
                if 'psi' in str(fck_obs.get('nota','')): notas.append('fck convertido de psi')
                ws_meta.cell(row=meta_row, column=14, value=' | '.join(notas))
                meta_row += 1
                cnt['meta_linhas'] += 1

            if cis_ok:
                if len(cis_ok) == 1:
                    ci_R = cis_ok[0][2]; ci_t_S = cis_ok[0][3]
                else:
                    vals = [c[2] for c in cis_ok]
                    ci_R = round(sum(vals)/len(vals), 5)
                    tvals = [c[3] for c in cis_ok if c[3] is not None]
                    ci_t_S = round(sum(tvals)/len(tvals), 5) if tvals else None
                    coms.append('{} formulacoes na METANALISE; col R/S = media.'.format(len(cis_ok)))
                cnt['col_R'] += 1
                if ci_t_S is not None: cnt['col_S'] += 1
                if any('cubo' in str(c[0].get('nota','')) for c in cis_ok):
                    coms.append('fck convertido cubo->cilindro (fator 0,8) D2.')
                if any((c[0].get('idade') or 0) > 28 for c in cis_ok):
                    coms.append('fck idade > 28d (sinalizado D6).')

        if ci_R is None and not coms:
            if not em_list:
                coms.append('Emissao nao-extraivel pelo parser. Revisar manualmente.')
                atipicos.append((str(estudo), 'emissao nao parseada'))
            elif not fck_list:
                coms.append('fck nao-extraivel pelo parser. Revisar manualmente.')
                atipicos.append((str(estudo), 'fck nao parseado'))
            else:
                coms.append('Erro de pareamento. Revisar manualmente.')
                atipicos.append((str(estudo), 'erro pareamento'))

        ws_an.cell(row=ws_row, column=18, value=ci_R if ci_R is not None else '')
        ws_an.cell(row=ws_row, column=19, value=ci_t_S if ci_t_S is not None else '')
        com_final = ' '.join(coms) if coms else ''
        ws_an.cell(row=ws_row, column=20, value=com_final)
        if com_final:
            ws_an.cell(row=ws_row, column=20).alignment = Alignment(wrap_text=True, vertical='top')
            if ci_R is None: cnt['sem_ci_com_nota'] += 1

        log_rows.append({
            'linha_excel': ws_row, 'estudo': str(estudo),
            'fck_raw': str(fck_raw or '')[:200],
            'fck_parseado': str([{k: r[k] for k in ('fck','tipo_cp','idade','label')} for r in fck_list]),
            'emissoes_raw': str(em_raw or '')[:250],
            'emissoes_parseadas': str([{k: e.get(k) for k in ('valor','unidade','por_m3','ci_direto','label','flag_excluir')} for e in em_list]),
            'uf_raw': str(uf_raw or '')[:150], 'durabilidade_raw': str(durab or '')[:150],
            'vida_util_anos': vu if vu else '',
            'ci_R': ci_R if ci_R is not None else '',
            'ci_t_S': ci_t_S if ci_t_S is not None else '',
            'excluido': excl, 'motivo_exclusao': mot_excl or '',
            'comentarios_col_T': com_final[:400],
        })

    wb.save(SRC)
    print('OK Banco salvo')
    if log_rows:
        with open(LOG_CSV, 'w', newline='', encoding='utf-8-sig') as f:
            w = csv.DictWriter(f, fieldnames=list(log_rows[0].keys()))
            w.writeheader()
            w.writerows(log_rows)
        print('OK Log auditoria: ' + LOG_CSV)
    print('\n' + '='*60)
    print('SUMARIO FINAL')
    print('='*60)
    for k, v in cnt.items():
        print('  {}: {}'.format(k, v))
    print('\nCASOS ATIPICOS PARA REVISAO MANUAL ({}):'.format(len(atipicos)))
    for est, mot in atipicos[:30]:
        print('  - {}: {}'.format(est, mot))
    if len(atipicos) > 30:
        print('  ... e mais {} -- ver col T e CSV de auditoria'.format(len(atipicos)-30))
    return cnt, atipicos


if __name__ == '__main__':
    processar()

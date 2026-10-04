"""Aplica a releitura dos 11 artigos com QG > 0 e sem intervalo gravado (2026-10-04).

Fontes: lotes_ic/resultado_icA.json e resultado_icB.json (subagentes) + leitura pessoal
das figuras em alta resolucao (recortes do shrink-drive-pdf.yml com first_page/last_page).
Regras:
  - IC so preenche celulas vazias (ic_*_valor); o tipo/n/citacao acompanham o valor;
  - SOBREPOR: medicoes pessoais que substituem as do lote;
  - correccoes de fck por leitura pessoal (rotulos/pixel em alta resolucao);
  - 3017 CUNNINGHAM: desdobramento da linha 259 nas 9 misturas da Fig.3b / Fig.9.
Correr da raiz do repo, depois de lotes_volnorma/aplicar.py, e a seguir recalcular_ci.py
e classificar_base_portland.py.
"""
import glob
import json
import re

import openpyxl

BASE = 'meta_analysis_base.xlsx'
C_MAT, C_FCK, C_IDADE, C_FONTE, C_OBS = 7, 8, 10, 12, 13
C_ICF, C_ICF_T, C_ICF_N, C_ICF_C = 26, 27, 28, 29
C_ICC, C_ICC_T, C_ICC_N, C_ICC_C = 30, 31, 32, 33
HOJE = '2026-10-04'


def clean(s):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(s))


def nota(ws, r, txt):
    obs = ws.cell(r, C_OBS).value
    ws.cell(r, C_OBS).value = clean(f'{obs} [{txt}]' if obs else f'[{txt}]')


# ---------------------------------------------------------------- tipos do lote -> tipos da base
TIPO = {'cov': 'cv_percentual', 'percentis': 'intervalo_interquartil'}

# ---------------------------------------------------------------- 3203 VACLAVIK (leitura pessoal)
# Figs.7 e 8 (pp.11-12) do recorte /printer: meia-barra de erro medida em pixel contra a grelha
# de cada painel (DP, 6 metades de prisma por mistura, Sec.2.3.1). Rotulos impressos em alta
# resolucao corrigem 8 fck que tinham sido lidos na versao /screen.
VAC_DP = {  # material (prefixo) -> DP
    'C-1-100': 1.8, 'C-1-90+10 escoria grossa': 1.05, 'C-1-80+20 escoria grossa': 0.6, 'C-1-70+30 escoria grossa': 0.7,
    'C-2-100': 1.45, 'C-2-90+10 escoria grossa': 1.15, 'C-2-80+20 escoria grossa': 0.5, 'C-2-70+30 escoria grossa': 1.3,
    'C-3-100': 0.7, 'C-3-90+10 escoria grossa': 0.65, 'C-3-80+20 escoria grossa': 0.55, 'C-3-70+30 escoria grossa': 0.25,
    'C-4-100': 1.4, 'C-4-90+10 escoria grossa': 0.5, 'C-4-80+20 escoria grossa': 0.8, 'C-4-70+30 escoria grossa': 1.2,
    'C-1-90+10 escoria fina': 1.05, 'C-1-80+20 escoria fina': 0.8, 'C-1-70+30 escoria fina': 0.55,
    'C-2-90+10 escoria fina': 1.75, 'C-2-80+20 escoria fina': 1.9, 'C-2-70+30 escoria fina': 2.2,
    'C-3-90+10 escoria fina': 0.95, 'C-3-80+20 escoria fina': 1.55, 'C-3-70+30 escoria fina': 1.65,
    'C-4-90+10 escoria fina': 1.1, 'C-4-80+20 escoria fina': 1.05, 'C-4-70+30 escoria fina': 1.25,
}
VAC_FCK = {  # material -> (gravado, rotulo impresso em alta resolucao)
    'C-3-70+30 escoria grossa': (26.5, 26.6),
    'C-1-90+10 escoria fina': (56.6, 56.8),
    'C-2-90+10 escoria fina': (53.5, 53.6),
    'C-2-70+30 escoria fina': (37.8, 37.9),
    'C-3-90+10 escoria fina': (47.9, 47.5),
    'C-3-70+30 escoria fina': (33.5, 34.3),
    'C-4-80+20 escoria fina': (42, 42.5),
    'C-4-70+30 escoria fina': (35.2, 36.2),
}

# ---------------------------------------------------------------- 4512 SKIBICKI (leitura pessoal)
# Fig.11b p.14 a 300 dpi: barras = CoV (Sec.3.4), meia-barra / media; R0 17,4% = maximo do texto (17,41%)
SKI_CV = {1480: (17.4, 'R0'), 1482: (7.2, 'R20'), 1485: (13.0, 'R50')}

# ---------------------------------------------------------------- 2999 RATHNARAJAN (leitura pessoal)
# Fig.2 (p.5) em recorte ampliado: as medias gravadas de 5 misturas nao batiam com a figura
# (FW = linha cheia, SW = tracejada). Os DP do lote ja estao emparelhados com as medias da figura.
RATH_FCK = {1905: (43, 44.8, 'CF30-FW'), 1906: (44, 42.0, 'CF30-SW'), 1912: (44, 45.7, 'CS30F15-FW'),
            1913: (46, 43.6, 'CS30F15-SW'), 1921: (45, 50.8, 'CS30M15-SW')}

# ---------------------------------------------------------------- 2714 KUL: desvio maximo global
KUL_TIPO = 'desvio_maximo_global'

# ---------------------------------------------------------------- 3685 MALAMI: IC de CO2 errado
MALAMI_LIMPAR = 1079

# ---------------------------------------------------------------- 3017 CUNNINGHAM: desdobramento
# Fig.3b (28 d, media +/- DP, n=5, cubos de 51 mm ASTM C109) e Fig.9 barras "Measured"
# (kg CO2-eq/m3, pixel: grelha 500/450/400/350, 6,25 px/kg). M = cinza moida; C = moida e tratada com CO2.
CUNN = [  # (mistura, fck, dp, co2)
    ('5%M-RHA', 54.82, 4.44, 459.5),  # linha 259 original (o fck 55 era desta mistura)
    ('CTRL', 48.89, 2.78, 482.2),
    ('5%M-WSA', 50.49, 3.83, 457.1),
    ('5%C-WSA', 45.55, 3.40, 456.3),
    ('15%M-WSA', 43.95, 6.36, 408.0),
    ('15%C-WSA', 26.54, 3.02, 405.3),
    ('5%C-RHA', 52.35, 2.53, 459.5),
    ('15%M-RHA', 53.95, 2.96, 414.4),
    ('15%C-RHA', 49.50, 2.47, 414.4),
]
CUNN_OBS = ('DESDOBRAMENTO 2026-10-04, leitura pessoal das figuras em alta resolucao: a linha legada juntava '
            'o fck da 5%M-RHA (~55) com o CO2 do CTRL (~480). Separada num traco por linha: fck aos 28 d = media '
            'da Fig.3b (barras de erro = 1 DP, n=5), CO2 = barra "Measured" da Fig.9 lida por pixel')


def desdobrar_cunningham(ws):
    src = 259
    assert str(ws.cell(src, 1).value) == '3017' and ws.cell(src, C_FCK).value == 55
    ncol = ws.max_column
    linhas = [src] + [ws.max_row + i for i in range(1, len(CUNN))]
    for r, (mix, fck, dp, co2) in zip(linhas, CUNN):
        if r != src:
            for c in range(1, ncol + 1):
                ws.cell(r, c).value = ws.cell(src, c).value
        ws.cell(r, C_MAT).value = f'Argamassa {mix} (cinza de biomassa: WSA palha de trigo / RHA casca de arroz; M moida, C moida e tratada com CO2)' if mix != 'CTRL' else 'Argamassa CTRL (100% cimento Portland, referencia)'
        ws.cell(r, C_FCK).value = fck
        ws.cell(r, C_IDADE).value = 28
        ws.cell(r, C_FONTE).value = f'Fig.3b p.8: {mix} 28 d {fck} +/- {dp} MPa (n=5)'
        ws.cell(r, C_OBS).value = f'[{CUNN_OBS}]'
        ws.cell(r, 14).value = co2
        ws.cell(r, 15).value = 'kg CO2-eq/m3'
        ws.cell(r, 16).value = 'Fig.9 p.11, barra "Measured" (pixel)'
        ws.cell(r, C_ICF).value = dp
        ws.cell(r, C_ICF_T).value = 'desvio_padrao'
        ws.cell(r, C_ICF_N).value = 5
        ws.cell(r, C_ICF_C).value = f'Fig.3b p.8: "Error bars show one standard deviation. (n = 5)"; {mix} {fck} +/- {dp}'
        ws.cell(r, 34).value = f'{co2} kg CO2-eq/m3 (leitura de grafico)'
        ws.cell(r, 35).value = 'Fig.9 p.11'
        ws.cell(r, 36).value = 'nenhuma - publicado em kg CO2-eq/m3 (lido por pixel)'
        ws.cell(r, 37).value = co2
        ws.cell(r, 38).value = None
        ws.cell(r, 39).value = None
        ws.cell(r, 40).value = None
        ws.cell(r, 41).value = None
    return set(linhas)


def main():
    res = {}
    for f in sorted(glob.glob('lotes_ic/resultado_ic*.json')):
        for a in json.load(open(f, encoding='utf-8')):
            for p in a['pares']:
                res[int(p['sheet_row'])] = (str(a['row_id']), p)
    wb = openpyxl.load_workbook(BASE)
    ws = wb['dados_fck']
    ignorar = desdobrar_cunningham(ws)

    n_f = n_c = 0
    for r, (rid, p) in sorted(res.items()):
        if r in ignorar or r in SKI_CV:
            continue
        assert str(ws.cell(r, 1).value) == rid, (r, rid)
        if p.get('ic_fck_valor') is not None and ws.cell(r, C_ICF).value in (None, ''):
            ws.cell(r, C_ICF).value = p['ic_fck_valor']
            ws.cell(r, C_ICF_T).value = TIPO.get(p['ic_fck_tipo'], p['ic_fck_tipo'])
            ws.cell(r, C_ICF_N).value = p.get('ic_fck_n')
            ws.cell(r, C_ICF_C).value = clean(f'[{HOJE}] {p.get("ic_fck_citacao") or ""} | {p.get("nota") or ""}')[:900]
            n_f += 1
        v = p.get('ic_co2_valor')
        if v is not None and ws.cell(r, C_ICC).value in (None, ''):
            ws.cell(r, C_ICC).value = v
            ws.cell(r, C_ICC_T).value = TIPO.get(p['ic_co2_tipo'], p['ic_co2_tipo'])
            ws.cell(r, C_ICC_N).value = p.get('ic_co2_n')
            ws.cell(r, C_ICC_C).value = clean(f'[{HOJE}] {p.get("ic_co2_citacao") or ""} | {p.get("nota") or ""}')[:900]
            n_c += 1
        if rid == '2714' and ws.cell(r, C_ICF).value is not None:
            ws.cell(r, C_ICF_T).value = KUL_TIPO

    # VACLAVIK
    n_vac = n_vac_fck = 0
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 1).value) != '3203':
            continue
        mat = str(ws.cell(r, C_MAT).value or '')
        chave = next((k for k in VAC_DP if mat.startswith(k)), None)
        if not chave:
            continue
        if chave in VAC_FCK:
            velho, novo = VAC_FCK[chave]
            assert ws.cell(r, C_FCK).value == velho, (r, chave, ws.cell(r, C_FCK).value)
            ws.cell(r, C_FCK).value = novo
            nota(ws, r, f'fck corrigido {HOJE}: rotulo impresso {novo} na Fig.{"7" if "grossa" in chave else "8"} em alta resolucao '
                        f'(recorte pp.11-12 do original); o {velho} tinha sido lido na versao /screen')
            n_vac_fck += 1
        if ws.cell(r, C_ICF).value in (None, ''):
            ws.cell(r, C_ICF).value = VAC_DP[chave]
            ws.cell(r, C_ICF_T).value = 'desvio_padrao'
            ws.cell(r, C_ICF_N).value = 6
            ws.cell(r, C_ICF_C).value = (f'[{HOJE}] Fig.{"7" if "grossa" in chave or chave.endswith("-100") else "8"} (pp.11-12), meia-barra de erro '
                                         f'medida em pixel contra a grelha do painel (recorte /printer); Sec.2.3.1: "6 samples were used '
                                         f'for determining compressive strength ... standard deviations ... depicted as error bars"')
            n_vac += 1

    # SKIBICKI
    for r, (cv, mix) in SKI_CV.items():
        assert str(ws.cell(r, 1).value) == '4512'
        if ws.cell(r, C_ICF).value in (None, ''):
            ws.cell(r, C_ICF).value = cv
            ws.cell(r, C_ICF_T).value = 'cv_percentual'
            ws.cell(r, C_ICF_C).value = (f'[{HOJE}] Fig.11b p.14 a 300 dpi: {mix} 28 d, meia-barra / media = {cv}% (barras = CoV, '
                                         f'Sec.3.4: "mean values ... and their respective CoV", 3,58-17,41%); n nao declarado')

    # RATHNARAJAN
    for r, (velho, novo, mix) in RATH_FCK.items():
        assert str(ws.cell(r, 1).value) == '2999' and ws.cell(r, C_FCK).value == velho, (r, ws.cell(r, C_FCK).value)
        ws.cell(r, C_FCK).value = novo
        nota(ws, r, f'fck corrigido {HOJE}, leitura pessoal da Fig.2 p.5 ampliada: {mix} aos 28 d = {novo} MPa '
                    f'(linha cheia = FW, tracejada = SW); o {velho} gravado nao batia com a figura')

    # MALAMI
    r = MALAMI_LIMPAR
    assert str(ws.cell(r, 1).value) == '3685' and ws.cell(r, C_ICC).value == 30.5
    for c in (C_ICC, C_ICC_T, C_ICC_N, C_ICC_C):
        ws.cell(r, c).value = None
    nota(ws, r, f'IC de CO2 removido {HOJE}: o 30,5 kg era o DP do ECO2 total a 100 anos da coluna PC (Tab.10), '
                f'nao da FA/35 (20,59) nem do CO2 A1-A3 por m3; a linha nao tem co2_kg_m3')

    wb.save(BASE)
    print(f'IC fck {n_f} | IC co2 {n_c} | VACLAVIK DP {n_vac}, fck {n_vac_fck} | CUNNINGHAM +{len(ignorar) - 1} linhas')


if __name__ == '__main__':
    main()

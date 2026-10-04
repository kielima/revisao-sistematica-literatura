"""Decisoes da Kie (2026-10-04) sobre os tres casos levantados na ronda do volume/IC.

  - 4321 ALMULHIM: fck = 'Factory Mix Fcr (28 Days MPa)' da Tabela 6 (38,27 ... 81,82), em vez
    dos rotulos C40-C80 (Fcr' aproximado);
  - 4416 ASDRUBALI: co2_kg_m3 = media (ponto branco) das caixas da Fig.5, em todas as classes;
  - 2833 OLIVEIRA: pares retirados (27,58 MPa da categoria '4 ksi' do BEES atribuido por
    extrapolacao as vigas A/B/C); sem outro par, o artigo sai em E2_resistencia (regra de 21/09).
Correr da raiz do repo e depois recalcular_ci.py.
"""
import re

import openpyxl

HOJE = '2026-10-04'
BASE, ART = 'meta_analysis_base.xlsx', 'articles_final.xlsx'

ALMULHIM = {1274: (40, 38.27), 1275: (50, 48.82), 1276: (60, 59.82), 1277: (70, 70.82), 1278: (80, 81.82)}
ASDRUBALI = {  # linha: (gravado, media da Fig.5)
    1336: (190, 185.0), 1337: (205, 205.9), 1338: (205, 204.5), 1339: (250, 237.0), 1340: (320, 319.6),
    1341: (325, 326.1), 1342: (360, 356.4), 1343: (385, 385.4), 1344: (250, 265.6), 1345: (290, 292.6),
}
OLIVEIRA = range(721, 728)


def clean(s):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(s))


def nota(ws, r, col, txt):
    v = ws.cell(r, col).value
    ws.cell(r, col).value = clean(f'{v} [{txt}]' if v else f'[{txt}]')


def main():
    wb = openpyxl.load_workbook(BASE)
    ws = wb['dados_fck']
    for r, (velho, novo) in ALMULHIM.items():
        assert str(ws.cell(r, 1).value) == '4321' and ws.cell(r, 8).value == velho, r
        ws.cell(r, 8).value = novo
        ws.cell(r, 10).value = 28
        nota(ws, r, 13, f'fck {velho} -> {novo} {HOJE} (decisao da Kie): Tabela 6, coluna "Factory Mix Fcr (28 Days MPa)"; '
                        f'o {velho} era o rotulo da mistura (Fcr\' aproximado) e a resistencia de projecto Fc e {velho - 10}')
    for r, (velho, media) in ASDRUBALI.items():
        assert str(ws.cell(r, 1).value) == '4416' and ws.cell(r, 37).value == velho, r
        ws.cell(r, 37).value = media
        ws.cell(r, 14).value = media
        ws.cell(r, 34).value = f'{media} kg CO2-eq/m3 (media da caixa, Fig.5 p.16)'
        nota(ws, r, 13, f'CO2 {velho} -> {media} {HOJE} (decisao da Kie): media (ponto branco) da caixa da Fig.5 lida por pixel; '
                        f'os valores antigos misturavam medianas e medias')
    for r in OLIVEIRA:
        assert str(ws.cell(r, 1).value) == '2833', r
        velho = ws.cell(r, 8).value
        ws.cell(r, 8).value = None
        nota(ws, r, 13, f'PAR RETIRADO {HOJE} (decisao da Kie): o fck {velho} e a categoria "4 ksi" do BEES 3.0, atribuida '
                        f'por extrapolacao as vigas A/B/C (tracos com CP IV-32 sem resistencia declarada)')
    wb.save(BASE)

    wa = openpyxl.load_workbook(ART)
    wt = wa.active
    r = 2833
    assert wt.cell(r, 20).value == 'Elegível' and wt.cell(r, 39).value == 'Aceito'
    wt.cell(r, 20).value = 'Rejeitado'
    wt.cell(r, 39).value = 'Rejeitado'
    wt.cell(r, 43).value = 2
    wt.cell(r, 44).value = 'E2_resistencia'
    wt.cell(r, 56).value = None
    nota(wt, r, 32, f'E2_resistencia Rejeitado {HOJE} (decisao da Kie): o unico fck (27,58 MPa) e a categoria "4 ksi" do BEES, '
                    f'atribuida por extrapolacao as vigas A/B/C do metodo proprio, sem resistencia declarada - sem par fck+CO2 '
                    f'proprio (regra E2 de 21/09)')
    wa.save(ART)
    print('ok')


if __name__ == '__main__':
    main()

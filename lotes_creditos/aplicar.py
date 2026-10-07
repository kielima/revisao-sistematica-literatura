"""CO2 sem desconto (bruto) e com desconto (creditos) - aplica a releitura de 2026-10-07 (44 lotes).

Pedido da Kie (2026-10-07): ver no painel os dois cenarios, com e sem desconto, como o filtro
«So cimento Portland». Decisoes F1-F3 em metanalise/DECISOES_METANALISE.md.

Escreve em meta_analysis_base.xlsx, aba dados_fck:
  col 37 co2_kg_m3                 -> passa a ser o BRUTO quando o artigo o da (antes alguns eram liquidos)
  col 44 co2_kg_m3_com_desconto    -> CO2 com os descontos que o proprio artigo aplica (vazio = sem desconto)
  col 45 co2_desconto_tipo         -> nenhum | carga_evitada; carbonatacao; biogenico; modulo_D; outro | indeterminado
  col 46 co2_desconto_citacao      -> onde estao os numeros, conta, derivado sim/nao
  col 47 co2_bruto_disponivel      -> sim | nao (o gravado tem desconto e o artigo nao permite obter o bruto)
Correr da raiz do repo; depois recalcular_ci.py, preparar_dados.py e a exportacao.
"""
import glob
import json
import re

import openpyxl

HOJE = '2026-10-07'
BASE = 'meta_analysis_base.xlsx'
COLS = {44: 'co2_kg_m3_com_desconto', 45: 'co2_desconto_tipo', 46: 'co2_desconto_citacao', 47: 'co2_bruto_disponivel'}

# Ajustes meus (por confirmar) sobre as respostas dos subagentes
AJUSTE_COM_DESCONTO = {
    594: (375.6, 'so a carbonatacao do proprio betao (-18,97 t); o albedo (-69,60 t) nao e credito do material e fica de fora'),
}


def clean(s):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(s))


def num(x):
    return x if isinstance(x, (int, float)) and not isinstance(x, bool) else None


def carregar():
    out = {}
    for f in sorted(glob.glob('lotes_creditos/resultado_lote*.json')):
        d = json.load(open(f, encoding='utf-8'))
        if isinstance(d, dict):
            d = d.get('estudos') or d.get('resultados') or next(v for v in d.values() if isinstance(v, list))
        for s in d:
            for p in s['pares']:
                out[int(p['sheet_row'])] = dict(p, row_id=s.get('row_id'), lote=f.rsplit('/', 1)[-1])
    return out


def main():
    res = carregar()
    wb = openpyxl.load_workbook(BASE)
    ws = wb['dados_fck']
    for c, nome in COLS.items():
        assert ws.cell(1, c).value in (None, nome), (c, ws.cell(1, c).value)
        ws.cell(1, c).value = nome
    n_bruto = n_desc = n_sem_bruto = 0
    for r, p in sorted(res.items()):
        assert str(ws.cell(r, 1).value) == str(p['row_id']), r
        gravado = num(ws.cell(r, 37).value)
        declara = str(p.get('artigo_declara_desconto') or '').lower()
        grav_e = str(p.get('gravado_e') or '').lower()
        bruto, com = num(p.get('co2_bruto')), num(p.get('co2_com_desconto'))
        tipos = p.get('tipos_desconto') or []
        if isinstance(tipos, str):
            tipos = [tipos]
        der = str(p.get('derivado') or '')
        cit = clean(f"{p.get('citacao') or ''} | derivado: {der} | {p.get('nota') or ''}")[:1500]
        if r in AJUSTE_COM_DESCONTO:
            com, motivo = AJUSTE_COM_DESCONTO[r]
            cit = clean(f'{cit} | AJUSTE {HOJE}: {motivo}')
        bruto_disp = 'sim'
        if grav_e in ('liquido', 'misto'):
            if com is None:
                com = gravado
            if bruto is not None and gravado is not None and abs(bruto - gravado) > 1e-6:
                ws.cell(r, 37).value = bruto
                ws.cell(r, 14).value = bruto
                v = ws.cell(r, 13).value
                ws.cell(r, 13).value = clean(f'{v} [CO2 {gravado} -> {bruto} {HOJE} (creditos): o gravado tinha desconto '
                                             f'({", ".join(tipos)}); passa a bruto e o liquido vai para co2_kg_m3_com_desconto]'
                                             if v else f'[CO2 {gravado} -> {bruto} {HOJE} (creditos)]')
                n_bruto += 1
            elif bruto is None:
                bruto_disp = 'nao'
                n_sem_bruto += 1
        elif grav_e == 'indeterminado' or declara == 'indeterminado':
            bruto_disp = 'nao' if declara != 'nao' and grav_e != 'bruto' else 'sim'
            n_sem_bruto += bruto_disp == 'nao'
        if declara == 'sim' and com is not None:
            ws.cell(r, 44).value = com
            n_desc += 1
        else:
            ws.cell(r, 44).value = None
        ws.cell(r, 45).value = ('; '.join(tipos) if tipos else 'sim (tipo nao indicado)') if declara == 'sim' else \
            ('indeterminado' if declara == 'indeterminado' else 'nenhum')
        ws.cell(r, 46).value = cit
        ws.cell(r, 47).value = bruto_disp
    wb.save(BASE)
    print(f'{len(res)} pares; {n_bruto} passaram de liquido a bruto; {n_desc} com valor com desconto; '
          f'{n_sem_bruto} sem bruto disponivel')


if __name__ == '__main__':
    main()
